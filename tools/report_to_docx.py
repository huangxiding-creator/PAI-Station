"""通用调研报告 Markdown → DOCX（复用 foundry/export_docx.py 中文字体排版基建）。

约定：首个 `# ` 为文档标题（进封面）；其后 `# ` 为章（Heading 1）、`## ` 为节
（Heading 2）；标题与首章之间的引言行进封面副题；其余交给 markdown-lite 渲染器
（h3/h4/表格/列表/引用/粗体）。自动生成点线目录与页码。

用法：
  python tools/report_to_docx.py <报告.md> <输出.docx>
"""

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from docx import Document

from paistation.foundry.export_docx import (  # 排版基建单一事实源
    GRAY, _center, _page_number, _render_markdown_lite, _setup, _toc_line,
)


def parse(md_text: str):
    """md → (标题, 封面副题行列表, 章列表)。章 = {title, children:[(h2|ml, ...)]}。"""
    title, cover, chapters = None, [], []
    cur_ch, buf = None, []

    def flush() -> None:
        nonlocal buf
        if buf and cur_ch is not None:
            cur_ch["children"].append(("ml", buf))
            buf = []

    for raw in md_text.splitlines():
        line = raw.rstrip()
        if line.startswith("# ") and not line.startswith("##"):
            text = line[2:].strip()
            if title is None:
                title = text
                continue
            flush()
            cur_ch = {"title": text, "children": []}
            chapters.append(cur_ch)
        elif line.startswith("## "):
            if cur_ch is None:
                cover.append(line[3:].strip())
            else:
                flush()
                cur_ch["children"].append(("h2", line[3:].strip()))
        elif title is not None and cur_ch is None:
            if line.strip():
                cover.append(line.lstrip("> ").strip())
        else:
            buf.append(line)
    flush()
    return title or "未命名报告", cover, chapters


def report_to_docx(md_path: str, out_path: str) -> str:
    title, cover, chapters = parse(Path(md_path).read_text(encoding="utf-8"))
    doc = Document()
    _setup(doc)

    _center(doc, title, size=22, bold=True, before_pt=150)
    for line in cover:
        _center(doc, line, size=11, color=GRAY)
    doc.add_page_break()

    doc.add_heading("目 录", level=1)
    for ch in chapters:
        _toc_line(doc, ch["title"])
        for kind, payload in ch["children"]:
            if kind == "h2":
                _toc_line(doc, payload, level=2)
    doc.add_page_break()

    for ch in chapters:
        doc.add_heading(ch["title"], level=1)
        for kind, payload in ch["children"]:
            if kind == "h2":
                doc.add_heading(payload, level=2)
            else:
                _render_markdown_lite(doc, "\n".join(payload))
    _page_number(doc)
    doc.save(out_path)
    return out_path


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    print(report_to_docx(sys.argv[1], sys.argv[2]))
