"""L1 提取：真格式夹具（在位库就地合成）+ 毒文件不崩批。"""
import email
import email.policy

from paistation.sense.localfiles.extract import (
    ExtractionError,
    extract,
    extract_eml,
    extract_pdf,
)
from paistation.sense.localfiles.triage import Triage


def _make_pdf(path):
    import pymupdf

    with pymupdf.open() as pdf:
        page = pdf.new_page()
        page.insert_text((72, 72), "PAI-Station 本地文件宇宙测试",
                         fontname="china-s")  # 内建简体字体（helv 写不了中文）
        pdf.save(str(path))


def _make_docx(path):
    import docx

    d = docx.Document()
    d.add_paragraph("项目周报：M2.6 连接器已落地")
    table = d.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "任务"
    table.rows[0].cells[1].text = "状态"
    d.save(str(path))


def _make_xlsx(path):
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "台账"
    ws.append(["项目", "进度"])
    ws.append(["LocalFiles", "P1"])
    wb.save(str(path))


def test_pdf_roundtrip(tmp_path):
    f = tmp_path / "报告.pdf"
    _make_pdf(f)
    doc = extract(str(f), "pdf", "pymupdf")
    assert "本地文件宇宙" in doc.text
    assert doc.parser_id == "pymupdf" and doc.parser_ver[0].isdigit()


def test_docx_roundtrip(tmp_path):
    f = tmp_path / "周报.docx"
    _make_docx(f)
    doc = extract(str(f), "word", "python-docx")
    assert "M2.6" in doc.text and "任务\t状态" in doc.text  # 段落+表格都进文本层


def test_xlsx_roundtrip(tmp_path):
    f = tmp_path / "台账.xlsx"
    _make_xlsx(f)
    doc = extract(str(f), "excel", "openpyxl")
    assert "## Sheet: 台账" in doc.text and "LocalFiles" in doc.text


def test_plaintext_and_eml(tmp_path):
    txt = tmp_path / "note.md"
    txt.write_text("# 标题\n内容", encoding="utf-8")
    doc = extract(str(txt), "text", "plaintext")
    assert doc.text.startswith("# 标题")

    eml = tmp_path / "mail.eml"
    msg = email.message.EmailMessage(policy=email.policy.SMTP)
    msg["Subject"] = "合同评审提醒"
    msg["From"] = "a@x.com"
    msg["To"] = "b@y.com"
    msg.set_content("请本周内回复评审意见。")
    eml.write_bytes(msg.as_bytes())
    doc = extract_eml(str(eml))
    assert doc.title == "合同评审提醒"
    assert "评审意见" in doc.text
    assert doc.meta["from"] == "a@x.com"


def test_poison_file_raises_not_crashes(tmp_path):
    fake = tmp_path / "毒.pdf"
    fake.write_bytes(b"not a real pdf but has .pdf suffix")
    try:
        extract(str(fake), "pdf", "pymupdf")
        raise AssertionError("应抛 ExtractionError")
    except ExtractionError:
        pass  # 单文件失败绝不崩批


def test_empty_pdf_detected_as_scan_candidate(tmp_path):
    import pymupdf

    f = tmp_path / "扫描件.pdf"
    with pymupdf.open() as pdf:  # 有页无文本层 → 留给 OCR 阶段
        pdf.new_page()
        pdf.save(str(f))
    try:
        extract_pdf(str(f))
        raise AssertionError("空 PDF 应抛 OCR 候选")
    except ExtractionError as exc:
        assert "OCR" in str(exc)


def test_triage_routes(tmp_path):
    t = Triage()
    assert t.classify("a/b/报告.pdf") == ("pdf", "pymupdf")
    assert t.classify("a/b/x.DOCX") == ("word", "python-docx")
    assert t.classify("a/b/邮件.msg") == ("email_msg", "extract_msg")
    assert t.classify("a/b/note.md") == ("text", "plaintext")
    assert t.classify("a/b/照片.png") == ("image", "metadata-only")
    # 后缀撒谎：无后缀但 %PDF 头 → 仍走 pdf
    f = tmp_path / "noext"
    f.write_bytes(b"%PDF-1.7 fake")
    assert t.classify(str(f)) == ("pdf", "pymupdf")


def test_metadata_only_never_reads_content(tmp_path):
    f = tmp_path / "图.png"
    f.write_bytes(b"\x89PNG garbage")
    doc = extract(str(f), "image", "metadata-only")
    assert doc.text == "" and doc.meta["size"] > 0


# ---------- OCR 兜底（09-17 扫描件收编） ----------

def _make_scan_pdf(path, text="平顶山市白龟湖 尾款支付 说明"):
    """合成扫描件：PIL 画字→PNG→pymupdf 嵌图页，零文本层。"""
    import pymupdf
    from PIL import Image, ImageDraw, ImageFont

    font = ImageFont.truetype("C:/Windows/Fonts/msyh.ttc", 36)
    img = Image.new("RGB", (1200, 400), "white")
    ImageDraw.Draw(img).text((40, 160), text, fill="black", font=font)
    png = str(path) + ".png"
    img.save(png)
    with pymupdf.open() as pdf:
        pdf.new_page(width=600, height=200).insert_image(
            pymupdf.Rect(0, 0, 600, 200), filename=png)
        pdf.save(str(path))


def test_squeeze_ocr_spaces():
    from paistation.sense.localfiles.ocr import squeeze_ocr_spaces

    raw = "平 顶 山 市 白 龟 湖 3691 ， 409 万 元 Total Cost Paid"
    assert squeeze_ocr_spaces(raw) == "平顶山市白龟湖 3691，409 万元 Total Cost Paid"


def test_ocr_fallback_marks_parser_ver(tmp_path, monkeypatch):
    import paistation.sense.localfiles.ocr as ocrmod

    _make_scan_pdf(tmp_path / "scan.pdf")
    monkeypatch.setattr(ocrmod, "ocr_pdf",
                        lambda p: "白龟湖尾款支付凭证 OCR 全文")
    doc = extract_pdf(str(tmp_path / "scan.pdf"))
    assert "OCR 全文" in doc.text
    assert doc.parser_ver.endswith("+winrt-ocr1")
    assert doc.meta["ocr"] == "winrt-ocr1"


def test_ocr_zero_text_keeps_failing(tmp_path, monkeypatch):
    import paistation.sense.localfiles.ocr as ocrmod

    _make_scan_pdf(tmp_path / "blank.pdf")
    monkeypatch.setattr(ocrmod, "ocr_pdf", lambda p: "  \n  ")
    try:
        extract_pdf(str(tmp_path / "blank.pdf"))
        raise AssertionError("应当失败")
    except ExtractionError as exc:
        assert "纯图" in str(exc)


def test_ocr_unavailable_falls_back_to_error(tmp_path, monkeypatch):
    import paistation.sense.localfiles.ocr as ocrmod

    _make_scan_pdf(tmp_path / "scan.pdf")
    def boom(_p):
        raise ocrmod.OcrUnavailable("缺语言包")
    monkeypatch.setattr(ocrmod, "ocr_pdf", boom)
    try:
        extract_pdf(str(tmp_path / "scan.pdf"))
        raise AssertionError("应当失败")
    except ExtractionError as exc:
        assert "OCR 不可用" in str(exc)


def test_winrt_ocr_end_to_end_real(tmp_path):
    """真机 WinRT OCR smoke（无语言包机器自动跳过）。"""
    import pytest

    from paistation.sense.localfiles.ocr import has_ocr, ocr_pdf
    if not has_ocr():
        pytest.skip("本机无 WinRT OCR 语言包")
    p = tmp_path / "real_scan.pdf"
    _make_scan_pdf(p)
    text = ocr_pdf(str(p))
    assert len(text.strip()) > 5  # 渲染清晰的雅黑 200dpi 必出字
    assert any(k in text for k in ("白龟湖", "龟湖", "尾款"))


def test_pdf_ocr_gets_longer_budget():
    from paistation.sense.localfiles.indexer import (
        EXTRACT_TIMEOUT_S, _job_timeout_s)

    assert _job_timeout_s("pdf") == EXTRACT_TIMEOUT_S * 4
    assert _job_timeout_s("word") == EXTRACT_TIMEOUT_S


def test_pv_compatible_ocr_suffix():
    from paistation.sense.localfiles.indexer import _pv_compatible

    assert _pv_compatible("1.26.4+winrt-ocr1", "1.26.4")
    assert _pv_compatible("1.26.4", "1.26.4")
    assert not _pv_compatible("1.26.5+winrt-ocr1", "1.26.4")


def test_cache_hit_accepts_ocr_suffix(tmp_path):
    from paistation.sense.localfiles.inventory import Inventory

    inv = Inventory(tmp_path / "inv.db")
    inv.apply_scan([{"path": "a.pdf", "size": 10, "mtime": 1000.0,
                     "secret": 0}])
    inv.mark_extracted("a.pdf", "pymupdf", "1.26.4+winrt-ocr1",
                       "deadbeef", "pdf")
    assert inv.cache_hit("a.pdf", 10, 1000.0, "pymupdf", "1.26.4")
    assert not inv.cache_hit("a.pdf", 10, 1000.0, "pymupdf", "1.26.5")
    inv.close()
