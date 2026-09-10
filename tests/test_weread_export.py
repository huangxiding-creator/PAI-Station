"""M2.5 铸造厂：微信读书导出层（MD/DOCX 专业排版）测试。"""
import os

from docx import Document

from paistation.forge.weread_export import (
    _anchor,
    _dedup_title_heading,
    _split_inline,
    to_docx,
    to_markdown,
)
from paistation.forge.weread_html import html_to_blocks

_HTML = ("<h1>小节标题</h1><p>网络协同<b>双螺旋</b>。</p>"
         "<blockquote>数据智能是引擎</blockquote>"
         "<img src=\"https://img/1.jpg\" alt=\"图解\">")


def _book():
    return {
        "bookId": "B1", "title": "智能商业", "author": "曾鸣",
        "format": "epub", "totalWords": 123456, "partial": False,
        "partialReasons": [], "intro": "曾鸣讲义。",
        "fetchedAt": "2026-09-10T12:00:00",
        "chapters": [
            {"chapterUid": 1, "title": "第一章 智能商业", "level": 1,
             "html": _HTML, "skipped": None},
            {"chapterUid": 2, "title": "第二章 数据智能", "level": 1,
             "html": "", "skipped": "free-limit"},
        ],
    }


def test_markdown_structure():
    md = to_markdown(_book())
    assert md.startswith("---\ntitle: 智能商业")
    assert "## 目录" in md
    assert "[第一章 智能商业](#第一章-智能商业)" in md
    assert "## 第一章 智能商业" in md                 # level1 章 → ##
    assert "### 小节标题" in md                        # 章内 h1 → ###
    assert "网络协同**双螺旋**。" in md
    assert "> 数据智能是引擎" in md
    assert "![图解](https://img/1.jpg)" in md
    assert "⚠️ 本章未提取：free-limit" in md


def test_markdown_image_mapper_localizes():
    md = to_markdown(_book(), image_mapper=lambda src: "images/1.jpg")
    assert "![图解](images/1.jpg)" in md


def test_markdown_partial_note():
    book = _book()
    book["partial"] = True
    book["partialReasons"] = ["第二章超出免费范围"]
    md = to_markdown(book)
    assert "**部分提取说明**：第二章超出免费范围" in md


def test_anchor_generation():
    assert _anchor("第一章 智能商业") == "第一章-智能商业"
    assert _anchor("AI 3.0") == "ai-30"


def test_dedup_title_heading_variants():
    """epub 章首复述章名：全等删除（连续多块）、带署名降级、不匹配原样。"""
    same = html_to_blocks("<h1>推荐序</h1><p>正文。</p>")[0]
    assert _dedup_title_heading(same, "推荐序") == same[1:]
    # 实测形态：para 复述 + h1 复述连续出现，全部剥除
    doubled = html_to_blocks("<p>版权信息</p><h1>版权信息</h1><p>书名：智能商业</p>")[0]
    assert _dedup_title_heading(doubled, "版权信息") == doubled[2:]
    mixed = html_to_blocks("<h1>推荐序<br/>马云|阿里巴巴集团创始人|</h1>")[0]
    out = _dedup_title_heading(mixed, "推荐序")
    assert out[0]["kind"] == "para" and "马云" in out[0]["text"]
    other = html_to_blocks("<h1>自序</h1>")[0]
    assert _dedup_title_heading(other, "推荐序") == other
    para_first = html_to_blocks("<p>无标题开头</p>")[0]
    assert _dedup_title_heading(para_first, "推荐序") == para_first
    # 全角空格等 Unicode 空白不影响判等（第一部分　智能商业）
    wsg = html_to_blocks("<h1>第一部分　智能商业</h1><p>内容</p>")[0]
    assert _dedup_title_heading(wsg, "第一部分 智能商业") == wsg[1:]


def test_split_inline():
    segs = _split_inline("普通**粗**与*斜*尾")
    assert segs == [("普通", False, False), ("粗", True, False),
                    ("与", False, False), ("斜", False, True),
                    ("尾", False, False)]


def test_docx_generates_with_layout(tmp_path):
    out = to_docx(_book(), str(tmp_path / "书.docx"))
    assert os.path.exists(out) and os.path.getsize(out) > 4000
    doc = Document(out)
    texts = [p.text for p in doc.paragraphs]
    # 封面：书名居中大字 + 免责行
    assert "智能商业" in texts and "仅供个人学习研究" in "".join(texts)
    # 目录页
    assert "目录" in texts and "第一章 智能商业" in texts
    # 正文：章标题 + 富文本（粗体独立 run）+ 图片占位
    assert "网络协同双螺旋。" in texts
    rich = [p for p in doc.paragraphs if p.text == "网络协同双螺旋。"][0]
    assert any(r.bold for r in rich.runs)
    assert any("〔图：图解〕" == p.text for p in doc.paragraphs)
    assert any("本章未提取：free-limit" == p.text for p in doc.paragraphs)
    # 页脚页码域存在
    footer_xml = doc.sections[0].footer.paragraphs[0]._p.xml
    assert "PAGE" in footer_xml


def test_docx_heading_styles_chinese_fonts(tmp_path):
    out = to_docx(_book(), str(tmp_path / "s.docx"))
    doc = Document(out)
    h1 = [p for p in doc.paragraphs if p.style.name == "Heading 1"][0]
    east = h1.style.element.rPr.rFonts.get(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
        "eastAsia")
    assert east == "微软雅黑"
