"""混沌学园 markdown → DOCX（metaso 专题库上传格式，递归批量）。

扫描 data/hundun/ 下所有课程目录（AI课程/、课程资料/<tab>/），
.md → .docx；已存在 .docx 跳过（幂等，供采集后增量转换）。
"""
import glob
import os
import sys

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(ROOT, "data", "hundun")


def add_para(doc, text, size=11, bold=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    return p


def convert(md_path: str) -> str:
    doc = Document()
    lines = open(md_path, encoding="utf-8").read().splitlines()
    for line in lines:
        s = line.strip()
        if not s:
            continue
        if s.startswith("# "):
            add_para(doc, s[2:], size=18, bold=True)
        elif s.startswith("## "):
            add_para(doc, s[3:], size=14, bold=True)
        elif s.startswith("- "):
            add_para(doc, "• " + s[2:])
        else:
            add_para(doc, s)
    out = os.path.splitext(md_path)[0] + ".docx"
    doc.save(out)
    return out


def main():
    # 递归收集课程 md（根级旧文件 + AI课程/ + 课程资料/<tab>/）；
    # 排除 _recon 档案与 skills/ 技能卡（非课程文稿）
    mds = [m for m in sorted(glob.glob(os.path.join(SRC, "**", "*.md"),
                                       recursive=True))
           if "_recon" not in m and f"{os.sep}skills{os.sep}" not in m]
    done = skipped = 0
    for md in mds:
        out = os.path.splitext(md)[0] + ".docx"
        if os.path.exists(out):
            skipped += 1
            continue
        convert(md)
        done += 1
        if done % 50 == 0:
            print(f"[docx] 已转换 {done}（跳过 {skipped}）", flush=True)
    print(f"[docx] 完成：转换 {done}，跳过已存在 {skipped}")


def convert_file(md: str) -> str:  # 兼容旧调用
    return convert(md)


if __name__ == "__main__":
    main()
