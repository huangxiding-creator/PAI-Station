# -*- coding: utf-8 -*-
"""终稿.md -> 终稿.docx (python-docx; 标题/段落/表格/粗体基础排版)."""
import re
import sys
from pathlib import Path

from docx import Document
from docx.shared import Pt

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
ROOT = Path(__file__).resolve().parent
src = (ROOT / "终稿.md").read_text(encoding="utf-8")

doc = Document()
style = doc.styles["Normal"]
style.font.name = "Times New Roman"
style.font.size = Pt(11)
from docx.oxml.ns import qn
style.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

lines = src.split("\n")
i = 0
while i < len(lines):
    ln = lines[i].rstrip()
    if not ln.strip() or ln.strip() == "---":
        i += 1
        continue
    # 表格块
    if ln.lstrip().startswith("|") and i + 1 < len(lines) and re.match(r"^\s*\|[-\s|:]+\|\s*$", lines[i + 1]):
        header = [c.strip() for c in ln.strip().strip("|").split("|")]
        rows = []
        j = i + 2
        while j < len(lines) and lines[j].lstrip().startswith("|"):
            rows.append([c.strip() for c in lines[j].strip().strip("|").split("|")])
            j += 1
        tb = doc.add_table(rows=1 + len(rows), cols=len(header))
        tb.style = "Light Grid Accent 1"
        for c, v in enumerate(header):
            tb.rows[0].cells[c].text = re.sub(r"\*\*(.+?)\*\*", r"\1", v)
        for r, row in enumerate(rows):
            for c in range(len(header)):
                v = row[c] if c < len(row) else ""
                tb.rows[r + 1].cells[c].text = re.sub(r"\*\*(.+?)\*\*", r"\1", v)
        doc.add_paragraph()
        i = j
        continue
    m = re.match(r"^(#{1,4})\s+(.*)$", ln)
    if m:
        lvl = min(4, len(m.group(1)))
        doc.add_heading(re.sub(r"\*\*(.+?)\*\*", r"\1", m.group(2)), level=lvl)
    elif re.match(r"^\s*[-*]\s+", ln):
        doc.add_paragraph(re.sub(r"^\s*[-*]\s+", "", ln), style="List Bullet")
    else:
        text = ln
        p = doc.add_paragraph()
        # 粗体分段
        for seg in re.split(r"(\*\*.+?\*\*)", text):
            if seg.startswith("**") and seg.endswith("**"):
                p.add_run(seg[2:-2]).bold = True
            elif seg:
                p.add_run(seg)
    i += 1

out = ROOT / "终稿.docx"
doc.save(str(out))
print("saved", out, out.stat().st_size, "bytes")
