#!/usr/bin/env python3
"""Recalculate formulas in a staged .xlsx/.xlsm via LibreOffice, in place.

This helper rewrites the workbook package. It is not an XML-identity or
preservation oracle and must not be run on the user's original workbook.

Usage: python recalc.py <workbook> [timeout_seconds=60]
Prints JSON: {"status": "success"|"errors_found", "total_formulas": N,
              "total_errors": N, "error_summary": {"#REF!": ["Sheet1!B10", ...]}}
or {"error": "..."} when nothing was recalculated (only this case exits non-zero).

Preflight: refuses to run when the workbook contains external-file references
(='[1]Sheet'!A1 style) — openpyxl re-saves strip their cached values and a
LibreOffice recalc would turn them into #NAME? and destroy the links.
Pass --force to accept that loss.
"""
import json
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

ERROR_TOKENS = ("#REF!", "#NAME?", "#VALUE!", "#DIV/0!", "#N/A", "#NULL!", "#NUM!")
MAX_LOCATIONS_PER_TYPE = 100


def fail(msg: str) -> None:
    print(json.dumps({"error": msg}, ensure_ascii=False))
    sys.exit(1)


def find_soffice() -> str:
    for cand in ("soffice", "libreoffice"):
        path = shutil.which(cand)
        if path:
            return path
    for path in (
        "/usr/bin/soffice",
        "/opt/libreoffice/program/soffice",
        "/Applications/LibreOffice.app/Contents/MacOS/soffice",
    ):
        if Path(path).exists():
            return path
    fail("LibreOffice (soffice) not found; cannot recalculate. "
         "Fall back to manual spot-checks and state clearly that no recalc ran.")
    raise SystemExit  # unreachable


def has_external_refs(path: Path) -> bool:
    import openpyxl  # deferred import so the error JSON stays clean

    wb = openpyxl.load_workbook(path, keep_vba=path.suffix.lower() == ".xlsm")
    pat = re.compile(r"='?\[\d+\]")
    try:
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str) and pat.search(cell.value):
                        return True
    finally:
        wb.close()
    return False


def main() -> None:
    args = [a for a in sys.argv[1:] if a != "--force"]
    force = "--force" in sys.argv
    if not args:
        fail("usage: recalc.py <workbook> [timeout_seconds] [--force]")
    src = Path(args[0]).resolve()
    timeout = int(args[1]) if len(args) > 1 else 60
    if not src.exists():
        fail(f"file not found: {src}")

    if not force and has_external_refs(src):
        fail("workbook contains external-file references; recalc would destroy "
             "them (#NAME?). Copy those cells' cached values out first, "
             "or rerun with --force to accept the loss.")

    soffice = find_soffice()
    with tempfile.TemporaryDirectory() as tmp:
        cmd = [
            soffice, "--headless", "--norestore",
            "--convert-to", src.suffix.lstrip("."),
            "--outdir", tmp, str(src),
        ]
        try:
            proc = subprocess.run(
                cmd, capture_output=True, text=True, timeout=timeout,
                env={"HOME": tmp, "PATH": "/usr/bin:/bin:/usr/local/bin"},
            )
        except subprocess.TimeoutExpired:
            fail(f"recalc timed out after {timeout}s")
        out = Path(tmp) / src.name
        if proc.returncode != 0 or not out.exists():
            fail(f"soffice failed: {proc.stderr.strip()[:500]}")
        shutil.move(str(out), str(src))

    import openpyxl

    wb_f = openpyxl.load_workbook(src, keep_vba=src.suffix.lower() == ".xlsm")
    wb_v = openpyxl.load_workbook(src, data_only=True)
    total_formulas = 0
    errors: dict[str, list[str]] = {}
    truncated = 0
    for ws_f, ws_v in zip(wb_f.worksheets, wb_v.worksheets):
        for row_f, row_v in zip(ws_f.iter_rows(), ws_v.iter_rows()):
            for cf, cv in zip(row_f, row_v):
                if isinstance(cf.value, str) and cf.value.startswith("="):
                    total_formulas += 1
                    val = cv.value
                    if isinstance(val, str) and val in ERROR_TOKENS:
                        locs = errors.setdefault(val, [])
                        if len(locs) < MAX_LOCATIONS_PER_TYPE:
                            locs.append(f"{ws_f.title}!{cf.coordinate}")
                        else:
                            truncated += 1
    wb_f.close()
    wb_v.close()

    total_errors = sum(len(v) for v in errors.values()) + truncated
    result = {
        "status": "errors_found" if total_errors else "success",
        "package_rewritten": True,
        "preservation_scope": "semantic checks must be performed separately",
        "total_formulas": total_formulas,
        "total_errors": total_errors,
        "error_summary": errors,
    }
    if truncated:
        result["locations_truncated"] = truncated
    print(json.dumps(result, ensure_ascii=False))


if __name__ == "__main__":
    main()
