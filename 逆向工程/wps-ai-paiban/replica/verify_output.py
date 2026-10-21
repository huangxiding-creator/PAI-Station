#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""AI排版复刻件输出实测验证器（数字零造假：全部指标直接从输出文档对象实测；
clear_span_props 后属性落在样式表，故字号/对齐/缩进按 run直接格式→样式链 的生效值判定）。"""
import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

FAILS = []


def check(name, cond, detail=""):
    tag = "PASS" if cond else "FAIL"
    print(f"[{tag}] {name}" + (f" — {detail}" if detail else ""))
    if not cond:
        FAILS.append(name)


def eff_size(p):
    if p.runs:
        for r in p.runs:
            if r.font.size is not None:
                return r.font.size.pt
    st = p.style
    seen = set()
    while st is not None and id(st) not in seen:
        seen.add(id(st))
        if st.font is not None and st.font.size is not None:
            return st.font.size.pt
        st = st.base_style
    return None


def eff_alignment(p):
    if p.alignment is not None:
        return str(p.alignment)
    st = p.style
    seen = set()
    while st is not None and id(st) not in seen:
        seen.add(id(st))
        if st.paragraph_format is not None and st.paragraph_format.alignment is not None:
            return str(st.paragraph_format.alignment)
        st = st.base_style
    return ""


def eff_ind(p):
    """沿 直接pPr→样式链 找 w:ind，返回元素或 None。"""
    pPr = p._p.pPr
    if pPr is not None and pPr.find(qn("w:ind")) is not None:
        return pPr.find(qn("w:ind"))
    st = p.style
    seen = set()
    while st is not None and id(st) not in seen:
        seen.add(id(st))
        el = st.element.find(qn("w:pPr"))
        if el is not None and el.find(qn("w:ind")) is not None:
            return el.find(qn("w:ind"))
        st = st.base_style
    return None


def main(path):
    doc = Document(path)
    body = doc.paragraphs

    # 1) 空段清零（不含 TOC 域段——其含域字符与占位文本）
    empties = [p for p in body if not p.text.strip() and p._p.find(".//" + qn("w:fldChar")) is None]
    check("空段清零", len(empties) == 0, f"残留 {len(empties)} 段")

    # 2) 主标题：独立样式、居中、18pt
    title = body[0]
    check("主标题挂【美化文档标题】", title.style.name == "美化文档标题", title.style.name)
    check("主标题居中", "CENTER" in eff_alignment(title), eff_alignment(title))
    check("主标题 18pt", eff_size(title) == 18.0, f"{eff_size(title)}pt")

    # 3) 一级标题：条款式（一、二、三）→ 全部左对齐【美化标题 1】（full-beauty.md 硬约束）
    h1s = [p for p in body if p.style.name in ("美化标题 1", "美化标题 1（居中）")]
    check("一级标题 ≥3 段", len(h1s) >= 3, f"{len(h1s)} 段")
    check("一级标题全部左对齐（存在条款式）", all(p.style.name == "美化标题 1" for p in h1s))
    check("一级标题 14pt", all(eff_size(p) == 14.0 for p in h1s), f"{[eff_size(p) for p in h1s]}")

    # 4) 正文：【美化正文】、两端对齐、首行缩进2字符(firstLineChars=200)、12pt
    bodies = [p for p in body if p.style.name == "美化正文" and p.text.strip()]
    check("正文段 ≥8 段", len(bodies) >= 8, f"{len(bodies)} 段")
    check("正文全部两端对齐", all("JUSTIFY" in eff_alignment(p) for p in bodies))
    ok_ind = all((lambda ind: ind is not None and ind.get(qn("w:firstLineChars")) == "200")(eff_ind(p))
                 for p in bodies)
    check("正文首行缩进 2 字符(firstLineChars=200)", ok_ind)
    check("正文 12pt", all(eff_size(p) == 12.0 for p in bodies), f"{sorted(set(eff_size(p) for p in bodies))}")

    # 5) 段首全角空格清零
    spaced = [p for p in body if p.text.strip() and p.text[:1] in ("　", " ", "\t")]
    check("段首空格缩进清零", len(spaced) == 0, f"残留 {len(spaced)} 段")

    # 6) 落款右对齐 + 首段段前距
    sig = [p for p in body if p.style.name in ("美化尾注", "美化尾注（首行）")]
    check("落款挂美化尾注 ≥2 段", len(sig) >= 2, f"{len(sig)} 段")
    check("落款右对齐", all("RIGHT" in eff_alignment(p) for p in sig))
    lead = next((p for p in sig if p.style.name == "美化尾注（首行）"), None)
    lead_before = None
    if lead is not None and lead.style.paragraph_format.space_before is not None:
        lead_before = lead.style.paragraph_format.space_before.pt
    check("落款首段段前 ≥24pt", lead_before is not None and lead_before >= 24, f"{lead_before}pt")

    # 7) 表格三线表
    if doc.tables:
        tbl = doc.tables[0]
        tb = tbl._tbl.tblPr.find(qn("w:tblBorders"))
        top = tb.find(qn("w:top")) if tb is not None else None
        insideV = tb.find(qn("w:insideV")) if tb is not None else None
        check("表格顶线 1.5pt", top is not None and top.get(qn("w:val")) == "single" and top.get(qn("w:sz")) == "12",
              f"val={top.get(qn('w:val')) if top is not None else None} sz={top.get(qn('w:sz')) if top is not None else None}")
        check("表格无竖线", insideV is not None and insideV.get(qn("w:val")) == "none")
        c0 = tbl.rows[0].cells[0]
        tcb = c0._tc.tcPr.find(qn("w:tcBorders")) if c0._tc.tcPr is not None else None
        bot = tcb.find(qn("w:bottom")) if tcb is not None else None
        check("表头底线 0.75pt", bot is not None and bot.get(qn("w:sz")) == "6")
        check("表头加粗", all(r.font.bold for p in c0.paragraphs for r in p.runs if p.text.strip()))

    # 8) 页脚含 PAGE 域
    xml = doc.sections[0].footer._element.xml
    check("页脚含 PAGE 域", "PAGE" in xml and "fldChar" in xml)

    # 9) TOC 域存在（1~3 级）
    has_toc = any("TOC" in (it.text or "") for p in body for it in p._p.iter(qn("w:instrText")))
    check("TOC 目录域已插入", has_toc)

    print()
    if FAILS:
        print(f"验证失败 {len(FAILS)} 项: {FAILS}")
        return 1
    print("全部验证通过 ✔")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1] if len(sys.argv) > 1 else str(Path(__file__).with_name("sample_messy_排版.docx"))))
