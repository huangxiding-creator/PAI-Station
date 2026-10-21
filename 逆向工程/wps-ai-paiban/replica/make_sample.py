#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""生成一份典型的"乱稿"docx，用于测试 AI排版复刻件。

刻意埋入的排版病灶（覆盖 WPS 规则库描述的常见问题）：
  - 大标题左对齐、字号 22pt 且用仿宋
  - 标题体系混用：第一章 / 一、 / （一） / 1. / 1.1
  - 正文三档字号混用（14pt / 15pt / 12pt）
  - 一半正文用全角空格硬缩进，一半无缩进
  - 3 个连续空段
  - 个别正文段误居中
  - 表格默认全边框、表头无强调
  - 落款左对齐（应右对齐）
"""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt

OUT = Path(__file__).with_name("sample_messy.docx")


def para(doc, text, size=12, bold=False, font="宋体", align=None, indent_spaces=0):
    p = doc.add_paragraph()
    r = p.add_run("　" * indent_spaces + text if indent_spaces else text)
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.name = "Times New Roman"
    r._r.rPr.rFonts.set(qn("w:eastAsia"), font)
    if align:
        p.alignment = align
    return p


def main():
    doc = Document()
    para(doc, "工程项目管理数字化转型白皮书（2026）", size=22, bold=True, font="仿宋")   # 病灶：大标题不居中

    para(doc, "", size=12)
    para(doc, "", size=12)

    para(doc, "第一章 总体形势", size=15, bold=True, font="黑体")
    para(doc, "近年来，工程建设行业加速向数字化、智能化方向演进，项目管理模式正在经历深刻变革。本报告围绕项目管理的核心场景，系统梳理数字化转型的现状、路径与关键抓手。", size=14, indent_spaces=2)
    para(doc, "从行业实践看，领先企业已经将进度、质量、安全、成本等核心要素纳入统一数据平台管理，实现了全过程可视与可追溯。", size=15)
    para(doc, "但整体上，中小企业的数字化能力仍然薄弱，工具碎片化、数据孤岛等问题突出。", size=12, indent_spaces=2)

    para(doc, "")
    para(doc, "")
    para(doc, "")

    para(doc, "一、转型的三大驱动力", size=15, bold=True, font="黑体")                    # 与「第一章」混用的条款式 h1
    para(doc, "（一）政策驱动", size=13, bold=True, font="楷体")
    para(doc, "各级主管部门陆续出台数字化转型指导意见，明确阶段性目标与评价标准，形成持续的政策推力。", size=12, indent_spaces=2)
    para(doc, "1. 标准体系逐步完善", size=12, bold=True)
    para(doc, "1.1 数据标准", size=12, bold=True)
    para(doc, "行业级数据标准与接口规范的发布，为跨系统协同奠定了基础。", size=14)
    para(doc, "（二）市场驱动", size=13, bold=True, font="楷体")
    para(doc, "业主方对交付确定性要求提高，倒逼承包方以数字化手段压缩管理损耗。", size=12, indent_spaces=2)
    para(doc, "本段是误居中的正文，用于测试对齐归一。", size=12, align=WD_ALIGN_PARAGRAPH.CENTER)

    para(doc, "二、实施路径", size=15, bold=True, font="黑体")
    para(doc, "实施路径应遵循「先试点、再推广、后固化」的三步走策略，优先在进度管理与安全管理两个场景取得可量化的收益。", size=15, indent_spaces=2)
    para(doc, "下表给出两个场景的对比：", size=12)

    tbl = doc.add_table(rows=3, cols=3)
    tbl.style = "Table Grid"
    for j, h in enumerate(("场景", "核心指标", "预期收益")):
        tbl.rows[0].cells[j].text = h
    data = [("进度管理", "计划符合率", "+12%"), ("安全管理", "隐患闭环率", "+18%")]
    for i, row in enumerate(data, start=1):
        for j, v in enumerate(row):
            tbl.rows[i].cells[j].text = v

    para(doc, "三、结论", size=15, bold=True, font="黑体")
    para(doc, "数字化转型不是工具采购，而是管理能力的系统性重构。应以数据贯通为主线，以场景收益为标尺，稳步推进。", size=14, indent_spaces=2)

    para(doc, "")
    para(doc, "某工程建设咨询研究中心", size=12)                                            # 病灶：落款左对齐
    para(doc, "2026年10月21日", size=12)

    doc.save(OUT)
    print(f"[样张已生成] {OUT}")


if __name__ == "__main__":
    main()
