# -*- coding: utf-8 -*-
"""PAI-Station 提案 Word 转换器：Markdown → 专业排版 docx（封面/目录/样式/表格）"""
import re, sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
ACCENT = RGBColor(0xC0, 0x39, 0x2B)

def set_font(run, name="微软雅黑", size=11, color=None, bold=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name)
    run.font.size = Pt(size)
    if color: run.font.color.rgb = color
    run.font.bold = bold

def add_rich(par, text, size=11, name="微软雅黑"):
    for seg in re.split(r'(\*\*.+?\*\*)', text):
        if not seg: continue
        if seg.startswith('**') and seg.endswith('**'):
            set_font(par.add_run(seg[2:-2]), name, size, NAVY, bold=True)
        else:
            set_font(par.add_run(seg), name, size)

def add_table(doc, rows):
    cells = [ [c.strip() for c in r.strip().strip('|').split('|')] for r in rows ]
    cells = [r for r in cells if not all(re.fullmatch(r':?-{2,}:?', c or '---') for c in r)]
    if not cells: return
    ncol = max(len(r) for r in cells)
    t = doc.add_table(rows=len(cells), cols=ncol)
    t.style = 'Table Grid'; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(cells):
        for j in range(ncol):
            txt = row[j] if j < len(row) else ''
            cell = t.cell(i, j)
            cell.paragraphs[0].text = ''
            add_rich(cell.paragraphs[0], txt.replace('**',''), size=9)
            if i == 0:
                sh = OxmlElement('w:shd'); sh.set(qn('w:fill'), '1F3A5F')
                cell._tc.get_or_add_tcPr().append(sh)
                for r in cell.paragraphs[0].runs: r.font.color.rgb = RGBColor(0xFF,0xFF,0xFF); r.font.bold = True

def add_toc(doc):
    p = doc.add_paragraph()
    fld = OxmlElement('w:fldSimple')
    fld.set(qn('w:instr'), r'TOC \o "1-2" \h \z \u')
    run = OxmlElement('w:r'); t = OxmlElement('w:t')
    t.text = "（右键此处 → 更新域，生成目录）"
    run.append(t); fld.append(run); p._p.append(fld)

def convert(md_path, doc, first=False):
    lines = open(md_path, encoding='utf-8').read().splitlines()
    i, in_code = 0, False
    while i < len(lines):
        ln = lines[i]
        if ln.strip().startswith('```'):
            in_code = not in_code
            if in_code:
                buf = []
                i += 1
                while i < len(lines) and not lines[i].strip().startswith('```'):
                    buf.append(lines[i]); i += 1
                for b in buf:
                    p = doc.add_paragraph()
                    p.paragraph_format.space_after = Pt(0)
                    set_font(p.add_run(b), "Consolas", 8.5)
                in_code = False
            i += 1; continue
        if ln.strip().startswith('|') and i+1 < len(lines) and lines[i+1].strip().startswith('|'):
            rows = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                rows.append(lines[i]); i += 1
            add_table(doc, rows); continue
        s = ln.strip()
        if not s or s.startswith('>'):
            if s.startswith('>'):
                p = doc.add_paragraph()
                add_rich(p, s.lstrip('> ').replace('**',''), size=9)
                for r in p.runs: r.font.color.rgb = RGBColor(0x66,0x66,0x66)
            i += 1; continue
        m = re.match(r'^(#{1,4})\s+(.*)', s)
        if m:
            lvl, txt = len(m.group(1)), m.group(2)
            h = doc.add_heading('', level=min(lvl,3))
            set_font(h.add_run(txt), "微软雅黑", {1:18,2:14,3:12}.get(lvl,11), NAVY, bold=True)
            i += 1; continue
        if re.fullmatch(r'-{3,}', s): i += 1; continue
        if s.startswith('- ') or s.startswith('* '):
            p = doc.add_paragraph(style='List Bullet')
            add_rich(p, s[2:]); i += 1; continue
        if s.startswith(('1.','2.','3.','4.','5.','6.','7.','8.')) and s[1] == '.':
            p = doc.add_paragraph(style='List Number')
            add_rich(p, s[3:]); i += 1; continue
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.3
        add_rich(p, s)
        i += 1

doc = Document()
sec = doc.sections[0]
sec.top_margin = Cm(2.5); sec.bottom_margin = Cm(2.5)
sec.left_margin = Cm(2.8); sec.right_margin = Cm(2.8)

# 封面
for _ in range(6): doc.add_paragraph()
t = doc.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_font(t.add_run("PAI-Station"), "微软雅黑", 40, NAVY, bold=True)
st = doc.add_paragraph(); st.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_font(st.add_run("个人超级 AI 工作站 · 产品提案与商业计划书"), "微软雅黑", 16, RGBColor(0x44,0x44,0x44))
for _ in range(3): doc.add_paragraph()
for label, val in [("提案版本","V2.2（操作手册级 · 含智能复利引擎 + 三重角色红队质询应答）"),("提案日期","2026 年 9 月 6 日"),("状态","待批准（批准 / 修订 / 否决）"),("研发基线","350 项开源调研 · 4 大竞品官网级解剖 · 19 项缝合件 · 32 个思维模型（含曾鸣三部曲） · 16 章 + 19 附录 · 红队 13 问全落解")]:
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run(f"{label}：{val}"), "微软雅黑", 12, RGBColor(0x33,0x33,0x33))
doc.add_paragraph()
note = doc.add_paragraph(); note.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_font(note.add_run("比我的想法牛逼 100 倍 · 当前技术条件完全可实现 · 亿级财富路径"), "微软雅黑", 11, ACCENT, bold=True)
doc.add_page_break()

# 目录
h = doc.add_heading('', level=1); set_font(h.add_run("目  录"), "微软雅黑", 18, NAVY, bold=True)
add_toc(doc)
doc.add_page_break()

convert("PROPOSAL.md", doc)
doc.add_page_break()
convert("BUSINESS_MODEL.md", doc)

doc.save("PAI-Station提案.docx")
print("saved PAI-Station提案.docx")
