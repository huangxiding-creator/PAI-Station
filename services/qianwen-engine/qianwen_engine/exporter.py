# -*- coding: utf-8 -*-
"""咨询问答导出：docx（python-docx）与 pdf（fpdf2 + 中文字体）。

设计口径：
- 与前端 md2blocks 同源的行级 Markdown 解析（标题/列表/表格/引用/分隔线/加粗）
- docx = 真 Word 结构（真表格、真列表样式）；pdf = A4 图纸风排版
- 依赖与字体都是惰性加载：缺依赖/缺字体 → 明确异常 → API 层转 503/400，绝不带崩主服务
"""
from __future__ import annotations

import base64
import io
import re
from pathlib import Path
from typing import Optional

from . import config


class FontMissing(Exception):
    """中文字体未配置（DATA_DIR/fonts/*.ttf）"""


class FormatNotSupported(Exception):
    """未知导出格式"""


# ── 行级 Markdown → 结构块（p/h/ul/ol/quote/table/hr） ──
def _blocks(md: str) -> list:
    lines = (md or "").replace("\r\n", "\n").split("\n")
    out: list = []
    i = 0
    while i < len(lines):
        ln = lines[i].rstrip()
        s = ln.strip()
        if not s:
            i += 1
            continue
        if s.startswith("#"):
            level = len(s) - len(s.lstrip("#"))
            out.append(("h", min(level, 4), s[level:].strip()))
            i += 1
            continue
        if s.startswith(">"):
            buf = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                buf.append(lines[i].strip().lstrip(">").strip())
                i += 1
            out.append(("quote", " ".join(x for x in buf if x)))
            continue
        if s.startswith("|") and i + 1 < len(lines) and re.match(r"^\|[\s:\-|]+\|?$", lines[i + 1].strip()):
            head = [c.strip() for c in s.strip("|").split("|")]
            i += 2
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            out.append(("table", head, rows))
            continue
        if s.startswith(("- ", "* ")):
            items = []
            while i < len(lines) and lines[i].strip().startswith(("- ", "* ")):
                items.append(lines[i].strip()[2:].strip())
                i += 1
            out.append(("ul", items))
            continue
        if re.match(r"^\d+[.、)]", s):
            items = []
            while i < len(lines):
                mm = re.match(r"^(\d+)[.、)]\s*(.*)$", lines[i].strip())
                if not mm:
                    break
                items.append(mm.group(2).strip())
                i += 1
            out.append(("ol", items))
            continue
        if re.match(r"^-{3,}$", s):
            out.append(("hr",))
            i += 1
            continue
        buf = [s]
        i += 1
        while i < len(lines):
            nxt = lines[i].strip()
            if (not nxt or re.match(r"^(#|>|\||-\s|\*\s|\d+[.、)])", nxt) or re.match(r"^-{3,}$", nxt)):
                break
            buf.append(nxt)
            i += 1
        out.append(("p", " ".join(buf)))
    return out


def _inline_runs(text: str) -> list:
    """**加粗** → [(text, bold), ...]；其余行内标记剥壳为纯文本。"""
    parts = re.split(r"(\*\*[^*]+\*\*)", text or "")
    runs = []
    for p in parts:
        if p.startswith("**") and p.endswith("**") and len(p) > 4:
            runs.append((p[2:-2], True))
        elif p:
            runs.append((re.sub(r"\*([^*]+)\*|`([^`]+)`", r"\1\2", p), False))
    return runs


def _plain(text: str) -> str:
    return "".join(v for v, _ in _inline_runs(text))


def _font_path() -> Optional[Path]:
    """中文字体查找：DATA_DIR/fonts/*.ttf 优先（次 *.otf，v0.7.4 海报 Noto pair），
    再扫系统常见位。"""
    fonts_dir = config.DATA_DIR / "fonts"
    if fonts_dir.is_dir():
        for pat in ("*.ttf", "*.otf"):
            for p in sorted(fonts_dir.glob(pat)):
                return p
    for cand in (
        Path("/usr/share/fonts/truetype/simhei/simhei.ttf"),
        Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"),
        Path("C:/Windows/Fonts/simhei.ttf"),
    ):
        if cand.is_file():
            return cand
    return None


def _doc_no(row: dict) -> str:
    d = (row.get("created_at") or "")[:10].replace("-", "")
    return (d[4:] or "0000") + "-" + str(row.get("id") or "")[:6].upper()


# ═══════════════ DOCX ═══════════════
def _build_docx(sections: list, title: str) -> bytes:
    """sections = [(question, md, citations, row), ...]（单条导出=1 节；批量=多节）。"""
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml.ns import qn
        from docx.shared import Pt, RGBColor
    except ImportError as exc:  # noqa: F401
        raise FormatNotSupported("服务器未安装 python-docx，Word 导出暂不可用") from exc

    NAVY = RGBColor(0x0D, 0x25, 0x47)
    ORANGE = RGBColor(0xE0, 0x7A, 0x0E)
    GRAY = RGBColor(0x7A, 0x84, 0x96)
    doc = Document()

    # 全局默认字体（Word 里 CJK 必须显式设 eastAsia）
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")

    def para(text_runs, size=11, bold=False, color=None, style=None,
             space_before=0, space_after=6, align=None):
        p = doc.add_paragraph(style=style)
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        if align is not None:
            p.alignment = align
        for val, b in text_runs:
            r = p.add_run(val)
            r.font.size = Pt(size)
            r.font.bold = bool(bold or b)
            r.font.name = "Calibri"
            r.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
            if color is not None:
                r.font.color.rgb = color
        return p

    # ── 抬头：图签风 ──
    para([(title, True)], size=17, color=NAVY, space_after=2)
    sub = " · 7×24小时工程智囊" if len(sections) == 1 else f" · 共 {len(sections)} 条咨询记录"
    para([(f"NO. {_doc_no(sections[0][3])}{sub}", False)], size=9, color=GRAY, space_after=4)

    for si, (question, md, citations, row) in enumerate(sections):
        if si > 0:
            para([("═" * 42, False)], size=9, color=GRAY, align=WD_ALIGN_PARAGRAPH.CENTER,
                 space_before=18, space_after=10)
            para([(f"NO. {_doc_no(row)}", False)], size=9, color=GRAY, space_after=6)
        para([("咨询问题", True)], size=11, color=ORANGE, space_before=6, space_after=2)
        para([(question or "", True)], size=12, color=NAVY, space_after=8)

        # ── 正文 ──
        for b in _blocks(md):
            kind = b[0]
            if kind == "h":
                level = b[1]
                size = 14 if level <= 1 else (12.5 if level == 2 else 11.5)
                para(_inline_runs(b[2]), size=size, bold=True, color=NAVY,
                     space_before=10, space_after=4)
            elif kind == "p":
                para(_inline_runs(b[1]))
            elif kind == "ul":
                for it in b[1]:
                    para(_inline_runs(it), style="List Bullet", space_after=2)
            elif kind == "ol":
                for it in b[1]:
                    para(_inline_runs(it), style="List Number", space_after=2)
            elif kind == "quote":
                para(_inline_runs(b[1]), style="Intense Quote")
            elif kind == "table":
                head, rows = b[1], b[2]
                t = doc.add_table(rows=1 + len(rows), cols=len(head))
                t.style = "Table Grid"
                for j, cell in enumerate(head):
                    c = t.rows[0].cells[j]
                    c.text = ""
                    r = c.paragraphs[0].add_run(_plain(cell))
                    r.font.bold = True
                    r.font.size = Pt(10.5)
                    r.font.color.rgb = NAVY
                for ri, rowv in enumerate(rows):
                    for j in range(len(head)):
                        c = t.rows[ri + 1].cells[j]
                        c.text = ""
                        r = c.paragraphs[0].add_run(_plain(rowv[j]) if j < len(rowv) else "")
                        r.font.size = Pt(10.5)
                doc.add_paragraph().paragraph_format.space_after = Pt(2)
            elif kind == "hr":
                para([("─" * 38, False)], size=9, color=GRAY, align=WD_ALIGN_PARAGRAPH.CENTER)

        # ── 依据来源 ──
        if citations:
            para([("依据来源 · BASIS OF ANSWER", True)], size=10, color=GRAY,
                 space_before=14, space_after=4)
            for c in citations:
                loc = f"（第 {c.get('loc')} 条）" if c.get("loc") else ""
                para([(f"[{c.get('n', '-')}] {c.get('source', '')}{loc}", False)],
                     size=9.5, color=GRAY, space_after=2)

    para([("由 总包AI顾问（AI 检索行业知识库生成）生成 · 仅供参考，不构成正式法律意见", False)],
         size=8.5, color=GRAY, space_before=16)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ═══════════════ PDF ═══════════════
def _build_pdf(sections: list, title: str) -> bytes:
    try:
        from fpdf import FPDF
    except ImportError as exc:
        raise FormatNotSupported("服务器未安装 fpdf2，PDF 导出暂不可用") from exc
    font = _font_path()
    if font is None:
        raise FontMissing("PDF 中文字体未配置（放置 TTF 到 data/qianwen/fonts/ 即启用）")

    NAVY = (13, 37, 71)
    ORANGE = (224, 122, 14)
    INK = (37, 45, 61)
    GRAY = (122, 132, 150)

    pdf = FPDF(format="A4")
    pdf.set_margins(18, 16, 18)
    pdf.set_auto_page_break(auto=True, margin=16)
    pdf.add_page()
    pdf.add_font("CN", "", str(font))
    pdf.add_font("CN", "B", str(font))

    # ── 抬头：蓝图题头带 ──
    pdf.set_fill_color(8, 26, 54)
    pdf.set_text_color(234, 243, 255)
    pdf.set_font("CN", "B", 15)
    pdf.cell(0, 11, " " + title, fill=True, new_x="LMARGIN", new_y="NEXT")
    pdf.set_text_color(*GRAY)
    pdf.set_font("CN", "", 9)
    sub = " · 7×24小时工程智囊" if len(sections) == 1 else f" · 共 {len(sections)} 条咨询记录"
    pdf.cell(0, 6, f" NO. {_doc_no(sections[0][3])}{sub}", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(3)

    epw = pdf.epw

    def _body(text: str, size=10.5, style="", color=INK, lh=5.4, indent=0.0):
        pdf.set_font("CN", style, size)
        pdf.set_text_color(*color)
        if indent:
            pdf.set_x(pdf.l_margin + indent)
        pdf.multi_cell(epw - indent, lh, text)

    for si, (question, md, citations, row) in enumerate(sections):
        if si > 0:  # 分节：新页 + 小题头
            pdf.add_page()
            pdf.set_text_color(*GRAY)
            pdf.set_font("CN", "", 9)
            pdf.cell(0, 6, f" NO. {_doc_no(row)} · 第 {si + 1}/{len(sections)} 条",
                     new_x="LMARGIN", new_y="NEXT")
            pdf.ln(2)

        # ── 咨询问题带 ──
        pdf.set_fill_color(240, 244, 251)
        pdf.set_text_color(*ORANGE)
        pdf.set_font("CN", "B", 10)
        pdf.cell(0, 6.5, " 咨询问题", fill=True, new_x="LMARGIN", new_y="NEXT")
        pdf.set_text_color(*NAVY)
        pdf.set_font("CN", "B", 12)
        pdf.multi_cell(0, 6.6, question or "", fill=True)
        pdf.ln(4)

        # ── 正文 ──
        for b in _blocks(md):
            kind = b[0]
            if kind == "h":
                level = b[1]
                pdf.ln(1.5)
                _body(_plain(b[2]), size=13 if level <= 1 else 11.5, style="B", color=NAVY, lh=6.2)
                pdf.ln(0.8)
            elif kind == "p":
                _body(_plain(b[1]))
                pdf.ln(1.2)
            elif kind == "ul":
                for it in b[1]:
                    _body("· " + _plain(it), lh=5.2, indent=2)
                    pdf.ln(0.4)
                pdf.ln(0.8)
            elif kind == "ol":
                for n, it in enumerate(b[1], 1):
                    _body(f"{n}.  {_plain(it)}", lh=5.2, indent=2)
                    pdf.ln(0.4)
                pdf.ln(0.8)
            elif kind == "quote":
                pdf.set_fill_color(240, 244, 251)
                pdf.set_font("CN", "", 10)
                pdf.set_text_color(57, 68, 90)
                pdf.set_x(pdf.l_margin + 3)
                pdf.multi_cell(epw - 3, 5.2, _plain(b[1]), fill=True)
                pdf.ln(1.5)
            elif kind == "table":
                # PDF v1：表头加粗行 + 数据行「列：值」 bullets（避免长单元格截断）
                head, rows = b[1], b[2]
                pdf.ln(0.6)
                _body(" | ".join(_plain(h) for h in head), size=10, style="B", color=NAVY, lh=5)
                for rowv in rows:
                    pairs = []
                    for j, h in enumerate(head):
                        v = _plain(rowv[j]) if j < len(rowv) else ""
                        pairs.append(f"{_plain(h)}：{v}" if len(head) > 1 else v)
                    _body("- " + "；".join(pairs), size=9.5, lh=5, indent=2)
                pdf.ln(1.0)
            elif kind == "hr":
                pdf.ln(1)
                pdf.set_draw_color(230, 225, 210)
                pdf.line(pdf.l_margin, pdf.get_y(), pdf.l_margin + epw, pdf.get_y())
                pdf.ln(2)

        # ── 依据来源 ──
        if citations:
            pdf.ln(3)
            pdf.set_draw_color(230, 225, 210)
            pdf.line(pdf.l_margin, pdf.get_y(), pdf.l_margin + epw, pdf.get_y())
            pdf.ln(2)
            _body("依据来源 · BASIS OF ANSWER", size=9.5, style="B", color=GRAY)
            for c in citations:
                loc = f"（第 {c.get('loc')} 条）" if c.get("loc") else ""
                _body(f"[{c.get('n', '-')}] {c.get('source', '')}{loc}", size=9, color=GRAY, indent=2)
        pdf.ln(2)

    pdf.ln(4)
    _body("由 总包AI顾问（AI 检索行业知识库生成）生成 · 仅供参考，不构成正式法律意见", size=8.5, color=GRAY)
    return bytes(pdf.output())


# ═══════════════ 入口 ═══════════════
def _sections(rows: list) -> list:
    out = []
    for row in rows:
        citations = row.get("citations") or []
        if isinstance(citations, str):
            try:
                import json
                citations = json.loads(citations)
            except ValueError:
                citations = []
        out.append((row.get("question") or "", row.get("answer_full") or "",
                    citations, row))
    return out


def _build_md(sections: list, title: str) -> bytes:
    lines = ["# " + title, ""]
    for si, (question, md, citations, row) in enumerate(sections, 1):
        if si > 1:
            lines += ["---", ""]
        lines += [f"## {si}. 咨询问题", "", question, "", "### 专业解答", "", md, ""]
        if citations:
            lines += ["#### 依据来源", ""]
            for c in citations:
                loc = f"（第 {c.get('loc')} 条）" if c.get("loc") else ""
                lines.append(f"- [{c.get('n', '-')}] {c.get('source', '')}{loc}")
            lines.append("")
    lines += ["---", "", "*由 总包AI顾问（AI 检索行业知识库生成）生成 · 仅供参考，不构成正式法律意见*", ""]
    return "\n".join(lines).encode("utf-8")


def build(fmt: str, row: dict) -> dict:
    """单条导出：fmt: docx|pdf → {filename, b64, bytes}（b64 供小程序直接落盘）。"""
    return _render(fmt, _sections([row]), "咨询问答", _doc_no(row))


def build_all(fmt: str, rows: list) -> dict:
    """v0.5.0 批量导出：全部咨询记录（docx/pdf/md）。"""
    first_no = _doc_no(rows[0]) if rows else "0000-000000"
    day = (rows[0].get("created_at") or "")[:10].replace("-", "") or first_no[:4]
    return _render(fmt, _sections(rows), "咨询档案", day)


def _render(fmt: str, sections: list, label: str, no: str) -> dict:
    fmt = (fmt or "").strip().lower()
    if fmt == "docx":
        data = _build_docx(sections, f"总包AI顾问 · {label}")
        ext = "docx"
    elif fmt == "pdf":
        data = _build_pdf(sections, f"总包AI顾问 · {label}")
        ext = "pdf"
    elif fmt == "md":
        data = _build_md(sections, f"总包AI顾问 · {label}")
        ext = "md"
    else:
        raise FormatNotSupported("仅支持 docx / pdf / md")
    return {
        "filename": f"总包AI顾问-{label}-{no}.{ext}",
        "b64": base64.b64encode(data).decode("ascii"),
        "bytes": len(data),
    }
