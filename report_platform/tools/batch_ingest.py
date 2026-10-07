# -*- coding: utf-8 -*-
"""batch_ingest — F:\\新建文件夹\\AI总包创新院 存量报告批量上架摄取器.

策划表驱动 (人工梳理, 不做模糊分类): 每条 = 目录定位键 + SKU + 系列 + 定价.
每目录选 canonical docx (>50KB 最大) + pdf (最大), python-docx 抽元数据
(章标题/字数/表格数), pypdf 抽页数 → work/catalog.json.

去重决策 (v1, 冲突目录入 backlog 不上架):
  - 中核工程 3 目录 (与旗舰 CNNC-EPC 同主体, 防自相残杀)
  - 同企业双版本取「全景洞察」系列 (寰球/山东电力咨询/上海市政/浙江省份)
  - 华东院独家 3 变体 (完整版无 docx) → 波次二
  - 国标 GB/T 50358 / 论文合集 / 马斯克指北 / 中间成果 → 不上架

用法: python tools/batch_ingest.py [--deep]
  --deep 重抽全部元数据 (默认: catalog.json 已有条目只补缺)
"""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
ROOT = Path(__file__).resolve().parents[1]
SRC = Path(r"F:\新建文件夹\AI总包创新院")
WORK = ROOT / "work"
CATALOG = WORK / "catalog.json"

# ---------------------------------------------------------- 策划表
# (定位键=目录名子串, sku, 系列, 定价)  系列: ent=企业全景 prov=省份市场
# topic=专题实战 excl=独家研究
PLAN: list[dict] = [
    # ---- 企业全景洞察 ¥698 ----
    {"key": "中国公路工程咨询集团", "sku": "ENT-01", "cat": "ent", "price": 698},
    {"key": "中国寰球工程有限公司EPC总承包业务全景洞察", "sku": "ENT-02", "cat": "ent", "price": 698},
    {"key": "中国石油工程建设有限公司", "sku": "ENT-03", "cat": "ent", "price": 698},
    {"key": "广西电力设计研究院", "sku": "ENT-04", "cat": "ent", "price": 698},
    {"key": "中石化广州工程", "sku": "ENT-05", "cat": "ent", "price": 698},
    {"key": "山东电力工程咨询院有限公司", "sku": "ENT-06", "cat": "ent", "price": 698},
    {"key": "铭扬工程设计", "sku": "ENT-07", "cat": "ent", "price": 698},
    {"key": "上海市政总院EPC总承包业务全景洞察", "sku": "ENT-08", "cat": "ent", "price": 698},
    {"key": "中亿国际设计集团", "sku": "ENT-09", "cat": "ent", "price": 698},
    {"key": "中冶南方EPC总承包业务研究报告", "sku": "ENT-10", "cat": "ent", "price": 698},
    {"key": "中南电力设计院EPC业务探秘", "sku": "ENT-11", "cat": "ent", "price": 698},
    {"key": "中国五环工程", "sku": "ENT-12", "cat": "ent", "price": 698},
    {"key": "中国天辰工程", "sku": "ENT-13", "cat": "ent", "price": 698},
    {"key": "中国建筑第二工程局", "sku": "ENT-14", "cat": "ent", "price": 698},
    {"key": "中国汽车工业工程", "sku": "ENT-15", "cat": "ent", "price": 698},
    {"key": "中国电建华东院EPC经营管理研究报告", "sku": "ENT-16", "cat": "ent", "price": 698},
    {"key": "中南勘测设计研究院", "sku": "ENT-17", "cat": "ent", "price": 698},
    {"key": "成都勘测设计研究院", "sku": "ENT-18", "cat": "ent", "price": 698},
    {"key": "昆明勘测设计研究院", "sku": "ENT-19", "cat": "ent", "price": 698},
    {"key": "中国联合工程", "sku": "ENT-20", "cat": "ent", "price": 698},
    {"key": "中城科泽", "sku": "ENT-21", "cat": "ent", "price": 698},
    {"key": "中建三局集团", "sku": "ENT-22", "cat": "ent", "price": 698},
    {"key": "中建五局", "sku": "ENT-23", "cat": "ent", "price": 698},
    {"key": "中建八局", "sku": "ENT-24", "cat": "ent", "price": 698},
    {"key": "中铁第四勘察设计院", "sku": "ENT-25", "cat": "ent", "price": 698},
    {"key": "华陆工程科技", "sku": "ENT-26", "cat": "ent", "price": 698},
    {"key": "湖南省建筑设计院", "sku": "ENT-27", "cat": "ent", "price": 698},
    # ---- 省份市场机会 ¥598 ----
    {"key": "云南省EPC", "sku": "PROV-01", "cat": "prov", "price": 598},
    {"key": "天津市EPC", "sku": "PROV-02", "cat": "prov", "price": 598},
    {"key": "宁夏回族自治区EPC", "sku": "PROV-03", "cat": "prov", "price": 598},
    {"key": "山东省EPC总承包市场机会", "sku": "PROV-04", "cat": "prov", "price": 598},
    {"key": "浙江省EPC总承包市场机会", "sku": "PROV-05", "cat": "prov", "price": 598},
    {"key": "海南省EPC", "sku": "PROV-06", "cat": "prov", "price": 598},
    {"key": "湖南省EPC", "sku": "PROV-07", "cat": "prov", "price": 598},
    {"key": "福建省EPC", "sku": "PROV-08", "cat": "prov", "price": 598},
    {"key": "贵州省EPC", "sku": "PROV-09", "cat": "prov", "price": 598},
    {"key": "黑龙江省EPC", "sku": "PROV-10", "cat": "prov", "price": 598},
    {"key": "新疆EPC", "sku": "PROV-11", "cat": "prov", "price": 598},
    {"key": "河北省EPC", "sku": "PROV-12", "cat": "prov", "price": 598},
    {"key": "四川省EPC", "sku": "PROV-13", "cat": "prov", "price": 598},
    {"key": "青海省EPC", "sku": "PROV-14", "cat": "prov", "price": 598},
    # ---- 专题实战 ¥498-598 ----
    {"key": "DeepSeek在EPC项目中的应用研究", "sku": "TOPIC-01", "cat": "topic", "price": 598},
    {"key": "EPC总承包项目经理能力手册", "sku": "TOPIC-02", "cat": "topic", "price": 598},
    {"key": "中国企业EPC业务创效30招式", "sku": "TOPIC-03", "cat": "topic", "price": 598},
    {"key": "储能EPC总承包全流程", "sku": "TOPIC-04", "cat": "topic", "price": 598},
    {"key": "风电EPC总承包全流程", "sku": "TOPIC-05", "cat": "topic", "price": 598},
    {"key": "杨房沟水电站EPC实践", "sku": "TOPIC-06", "cat": "topic", "price": 598},
    {"key": "核电工程EPC项目管理策划编制指南", "sku": "TOPIC-07", "cat": "topic", "price": 598},
    {"key": "水利工程EPC总承包模式实施效果", "sku": "TOPIC-08", "cat": "topic", "price": 598},
    {"key": "水利工程EPC项目管理策划编制指南", "sku": "TOPIC-09", "cat": "topic", "price": 598},
    {"key": "江西丰城电厂三期", "sku": "TOPIC-10", "cat": "topic", "price": 598},
    {"key": "雅鲁藏布江下游", "sku": "TOPIC-11", "cat": "topic", "price": 598},
    {"key": "从100篇论文看电建华东院", "sku": "TOPIC-12", "cat": "topic", "price": 498},
    {"key": "工程总承包招投标合规工作指南", "sku": "TOPIC-13", "cat": "topic", "price": 498},
    {"key": "企业EPC总承包培训深度洞察", "sku": "TOPIC-14", "cat": "topic", "price": 498},
]

# 散文件 (排队区/根区/独家区): 显式路径
FILES: list[dict] = [
    {"sku": "ENT-28", "cat": "ent", "price": 698,
     "title": "中国建筑第六工程局有限公司EPC总承包业务全景洞察报告",
     "docx": r"总包研报\00 排队处理研报\《中国建筑第六工程局有限公司EPC总承包业务全景洞察报告》.docx",
     "need_pdf": True},
    {"sku": "PROV-15", "cat": "prov", "price": 598,
     "title": "湖北省EPC总承包市场开发工作指南（2025年版）",
     "docx": r"总包研报\《湖北省EPC总承包市场开发工作指南（2025年版）》（总包创研院）.docx",
     "pdf": r"总包研报\《湖北省EPC总承包市场开发工作指南（2025年版）》（总包创研院）.pdf"},
    {"sku": "PROV-16", "cat": "prov", "price": 598,
     "title": "山西省EPC总承包市场深度研究报告",
     "docx": r"总包研报\00 排队处理研报\山西省EPC总承包市场深度研究报告.docx",
     "need_pdf": True},
    {"sku": "TOPIC-15", "cat": "topic", "price": 498,
     "title": "风电EPC项目管理策划工作指南",
     "docx": r"总包研报\《风电EPC项目管理策划工作指南》.docx",
     "need_pdf": True},
    {"sku": "TOPIC-16", "cat": "topic", "price": 498,
     "title": "辽宁省固定资产投资分析报告",
     "docx": r"总包研报\00 排队处理研报\辽宁省固定资产投资分析报告_最终版(1).docx",
     "need_pdf": True},
    {"sku": "TOPIC-17", "cat": "topic", "price": 498,
     "title": "工程人这样用AI：从入门到精通，成为AI时代的超级项目管理者",
     "docx": r"总包研报\00 排队处理研报\工程人这样用AI：从入门到精通，成为AI时代的超级项目管理者_V2.docx",
     "need_pdf": True},
    {"sku": "EXCL-01", "cat": "excl", "price": 698,
     "title": "中国铁设如何做EPC总承包？",
     "docx": r"独家研究报告\《中国铁设如何做EPC总承包？》260617.docx",
     "pdf": r"独家研究报告\《中国铁设如何做EPC总承包？》260617.pdf"},
    {"sku": "EXCL-02", "cat": "excl", "price": 698,
     "title": "浙江省水利工程商机研究报告",
     "docx": r"独家研究报告\《浙江省水利工程商机研究报告》总包创研院.docx",
     "pdf": r"独家研究报告\《浙江省水利工程商机研究报告》总包创研院.pdf"},
]

CAT_NAME = {"ent": "企业全景洞察", "prov": "省份市场机会",
            "topic": "专题实战", "excl": "独家研究"}


def _find_dir(key: str) -> Path | None:
    base = SRC / "总包研报"
    hits = [d for d in base.iterdir()
            if d.is_dir() and key in d.name and "排队" not in d.name]
    return hits[0] if len(hits) == 1 else None


def _pick(folder: Path, suffix: str, min_kb: int) -> Path | None:
    cands = [f for f in folder.rglob(f"*{suffix}")
             if f.stat().st_size >= min_kb * 1024]
    return max(cands, key=lambda f: f.stat().st_size) if cands else None


def _meta(docx: Path, pdf: Path | None) -> dict:
    from docx import Document
    doc = Document(str(docx))
    chapters, chars, tables = [], 0, 0
    h2: list[str] = []
    for p in doc.paragraphs:
        txt = p.text.strip()
        if not txt:
            continue
        chars += len(re.findall(r"[\u4e00-\u9fff]", txt))
        st = (p.style.name or "").lower()
        if ("heading 1" in st or st == "title") and len(txt) <= 60:
            chapters.append(txt)
        elif "heading 2" in st and len(txt) <= 60:
            h2.append(txt)
    tables = len(doc.tables)
    for t in doc.tables:                      # \u8868\u5185\u6587\u5b57\u8ba1\u5165\u5b57\u6570 (\u6b63\u6587\u6bb5\u6f0f\u8ba1\u515c\u5e95)
        for row in t.rows:
            for cell in row.cells:
                chars += len(re.findall(r"[\u4e00-\u9fff]", cell.text))
    if not chapters:                          # \u515c\u5e95\u2460: Heading2 \u5f53\u7ae0
        chapters = h2
    if not chapters:                          # \u515c\u5e95\u2461: \u300c\u7b2cX\u90e8\u5206/\u7bc7\u300d\u77ed\u884c
        chapters = [p.text.strip() for p in doc.paragraphs
                    if len(p.text.strip()) <= 40
                    and re.match(r"^\u7b2c[\u4e00\u4e8c\u4e09\u56db\u4e94\u516d\u4e03\u516b\u4e5d\u5341\d]+[\u90e8\u5206\u7bc7\u7ae0]",
                                 p.text.strip())]
    pages = 0
    if pdf and pdf.is_file():
        try:
            from pypdf import PdfReader
            pages = len(PdfReader(str(pdf)).pages)
        except Exception:
            pages = 0
    # 样板前置章 (法律声明/版权声明/前言/摘要…) 不进商品目录编号
    boiler = re.compile(
        r"^(法律声明|版权声明|免责声明|前言|摘要|内容摘要|目录|序言|序$|"
        r"概述|引言|报告说明|编制说明|声明$)")
    chapters = [c for c in chapters if not boiler.match(c)]
    return {"chapters": chapters[:60], "words": chars, "tables": tables,
            "pages": pages}


def main(argv: list[str]) -> int:
    deep = "--deep" in argv
    WORK.mkdir(exist_ok=True)
    old = json.loads(CATALOG.read_text(encoding="utf-8")) \
        if CATALOG.is_file() else {"entries": []}
    have = {e["sku"]: e for e in old.get("entries", [])}
    entries: list[dict] = []

    def one(sku: str, cat: str, price: int, title: str,
            docx: Path, pdf: Path | None) -> dict:
        if not deep and sku in have and have[sku].get("words"):
            return have[sku]
        m = _meta(docx, pdf)
        e = {"sku": sku, "cat": cat, "cat_name": CAT_NAME[cat], "price": price,
             "title": title, "docx": str(docx), "pdf": str(pdf) if pdf else "",
             **m}
        print(f"[ing] {sku} {title[:26]:<28} {len(m['chapters']):>2}章 "
              f"{m['words']/10000:5.1f}万字 {m['tables']:>3}表 "
              f"{m['pages']:>3}页")
        return e

    zb = SRC / "总包研报"
    for row in PLAN:
        d = _find_dir(row["key"])
        if not d:
            print(f"[!!] 定位失败: {row['key']}")
            continue
        title = re.sub(r"^\d+\s*|^[1-9]", "", d.name).strip("《》 ")
        docx = _pick(d, ".docx", 50)
        pdf = _pick(d, ".pdf", 100)
        if not docx:
            print(f"[!!] 无 docx: {d.name}")
            continue
        entries.append(one(row["sku"], row["cat"], row["price"],
                           title, docx, pdf))
    for row in FILES:
        docx = SRC / row["docx"]
        pdf = (SRC / row["pdf"]) if row.get("pdf") else None
        if not docx.is_file():
            print(f"[!!] 散文件缺: {row['docx']}")
            continue
        entries.append(one(row["sku"], row["cat"], row["price"],
                           row["title"], docx, pdf))

    entries.sort(key=lambda e: e["sku"])
    CATALOG.write_text(json.dumps({"entries": entries}, ensure_ascii=False,
                                  indent=1), encoding="utf-8")
    ok = [e for e in entries if e.get("words")]
    no_pdf = [e for e in entries if not e["pdf"]]
    print(f"\n[ingest] {len(entries)} 条入册 ({len(ok)} 有元数据, "
          f"{len(no_pdf)} 缺PDF待COM转换) → {CATALOG}")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
