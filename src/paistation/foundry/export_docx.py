"""方案 → DOCX 专业排版（M7.2，复用 weread_export 的中文字体排版铁律）。

渲染 markdown-lite 子集：h1-h4 / 段落 / **粗体** / 表格 / 列表 / 引用。
封面（标题+密度分）→ 目录（章节点线）→ 正文。
"""

import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
GRAY = RGBColor(0x66, 0x66, 0x66)
RED = RGBColor(0xC0, 0x39, 0x2B)


def plan_to_docx(plan: dict, out_path: str) -> str:
    """方案 dict → DOCX 文件，返回路径。"""
    doc = Document()
    _setup(doc)

    _center(doc, plan.get("title", "未命名方案"), size=24, bold=True, before_pt=120)
    score = plan.get("score", 0)
    _center(doc, f"思想密度分：{score}（{'通过密度闸门' if plan.get('passed') else '未达阈值'}）",
            size=12, color=GRAY)
    _center(doc, "PAI-Station 方案铸造厂出品 · 反馈即返钱 · 详见方案页", size=10, color=GRAY)
    doc.add_page_break()

    doc.add_heading("目录", level=1)
    for ch in plan.get("chapters", []):
        _toc_line(doc, ch.get("title", ""))
        for sec in ch.get("sections", []):
            _toc_line(doc, sec.get("title", ""), level=2)
    doc.add_page_break()

    for ch in plan.get("chapters", []):
        doc.add_heading(ch.get("title", ""), level=1)
        _note(doc, f"章框架：{ch.get('framework', '未标注')}")
        for sec in ch.get("sections", []):
            doc.add_heading(sec.get("title", ""), level=2)
            comps = "、".join(sec.get("components", [])) or "无"
            _note(doc, f"节框架：{sec.get('framework', '未标注')}｜九件套组件：{comps}")
            _render_markdown_lite(doc, sec.get("content", ""))
            if sec.get("degraded"):
                _note(doc, f"⚠ {sec.get('degrade_note', '降级')}", color=RED)
    _page_number(doc)
    doc.save(out_path)
    return out_path


# ---------- 排版基建（与 weread_export 同铁律） ----------

def _setup(doc) -> None:
    for section in doc.sections:
        section.page_width, section.page_height = Cm(21.0), Cm(29.7)
        section.top_margin = section.bottom_margin = Cm(2.54)
        section.left_margin = section.right_margin = Cm(3.18)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    for name, size in (("Heading 1", 16), ("Heading 2", 14), ("Heading 3", 12)):
        st = doc.styles[name]
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.bold = True
        st.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "微软雅黑")


def _center(doc, text: str, size=11, bold=False, color=None, before_pt=0) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if before_pt:
        p.paragraph_format.space_before = Pt(before_pt)
    _runs(p, text, size=size, bold=bold, color=color)


def _note(doc, text: str, color=GRAY) -> None:
    p = doc.add_paragraph()
    _runs(p, text, size=9.5, color=color, italic=True)


def _toc_line(doc, text: str, level=1) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0 if level == 1 else 0.6)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(15.0), WD_TAB_ALIGNMENT.RIGHT,
                                              WD_TAB_LEADER.DOTS)
    _runs(p, text, size=11 if level == 1 else 10,
          bold=level == 1)


def _runs(p, text: str, size=11, bold=False, color=None, italic=False) -> None:
    for seg in re.split(r"(\*\*.+?\*\*)", text):
        if not seg:
            continue
        if seg.startswith("**") and seg.endswith("**"):
            run = p.add_run(seg[2:-2])
            run.font.bold = True
        else:
            run = p.add_run(seg)
            run.font.bold = bold
        run.font.italic = italic
        run.font.name = "Calibri"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        run.font.size = Pt(size)
        if color:
            run.font.color.rgb = color


def _render_markdown_lite(doc, md_text: str) -> None:
    """渲染 markdown-lite：h3/h4、表格、列表、引用、段落（h1/h2 已被章/节占用）。"""
    lines = md_text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip():
            i += 1
            continue
        if line.startswith("|") and i + 1 < len(lines) and \
                re.match(r"^\|[\s:|-]+\|?$", lines[i + 1].strip()):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(lines[i])
                i += 1
            _table(doc, rows)
            continue
        if line.startswith("####"):
            doc.add_heading(line.lstrip("# ").strip(), level=4)
        elif line.startswith("###"):
            doc.add_heading(line.lstrip("# ").strip(), level=3)
        elif re.match(r"^[-*]\s+", line):
            p = doc.add_paragraph(style="List Bullet")
            _runs(p, re.sub(r"^[-*]\s+", "", line))
        elif re.match(r"^\d+[.、]\s*", line):
            p = doc.add_paragraph(style="List Number")
            _runs(p, re.sub(r"^\d+[.、]\s*", "", line))
        elif line.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.8)
            _runs(p, line.lstrip("> ").strip(), color=GRAY)
        else:
            p = doc.add_paragraph()
            _runs(p, line.strip())
        i += 1


def _table(doc, rows) -> None:
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    cells = [r for r in cells if not all(re.fullmatch(r":?-{2,}:?", c or "---")
                                         for c in r)]
    if not cells:
        return
    ncol = max(len(r) for r in cells)
    t = doc.add_table(rows=len(cells), cols=ncol)
    t.style = "Table Grid"
    for ri, row in enumerate(cells):
        for ci in range(ncol):
            txt = row[ci] if ci < len(row) else ""
            cell = t.cell(ri, ci)
            cell.paragraphs[0].text = ""
            _runs(cell.paragraphs[0], txt.replace("**", ""), size=9,
                  bold=ri == 0)
            if ri == 0:
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), "1F3A5F")
                cell._tc.get_or_add_tcPr().append(shade)
                for run in cell.paragraphs[0].runs:
                    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)


def _page_number(doc) -> None:
    p = doc.sections[0].footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    for tag, attr in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            fld = OxmlElement("w:fldChar")
            fld.set(qn("w:fldCharType"), tag)
            run._r.append(fld)
        else:
            instr = OxmlElement("w:instrText")
            instr.text = attr
            run._r.append(instr)
