"""L1 提取器：路由表 → 统一 Document 中间表示。

统一 IR（Tika 式）：{path, kind, title, text, meta, attachments,
parser_id, parser_ver}。parser_ver 取各库实际版本号（升级自动
失效提取缓存）。质量阶梯：born-digital fast 文本层优先；无文本层
扫描件（白龟湖支付凭证/影像卷类核心证据）当场 WinRT OCR 兜底
（见 ocr.py，09-17 落地）。
崩溃隔离：per-file 异常全捕 + 容量硬顶；原生段错误风险由索引层的
线程超时护栏兜（见 indexer）。
"""
from __future__ import annotations

import email
import logging
import os
import sys
from dataclasses import dataclass, field

_log = logging.getLogger("paistation.sense.localfiles.extract")

MAX_TEXT_BYTES = 4 << 20        # 纯文本读入上限 4MB
MAX_PDF_PAGES = 2000
MAX_EXCEL_CELLS = 50_000


class ExtractionError(Exception):
    """单文件提取失败（毒文件常态化，绝不崩批）。"""


@dataclass
class Document:
    """统一中间表示：文本层 + 元数据 + 附件清单。"""
    path: str
    kind: str
    parser_id: str = ""
    parser_ver: str = ""
    title: str = ""
    text: str = ""
    meta: dict = field(default_factory=dict)
    attachments: list[str] = field(default_factory=list)

    @property
    def ok(self) -> bool:
        return bool(self.text or self.title)


# ---------- 各路由解析器 ----------

def _ver(mod) -> str:
    return str(getattr(mod, "__version__", getattr(mod, "VERSION", "?")))


def extract_pdf(path: str) -> Document:
    import pymupdf  # 1.28+；不再用已废弃的 fitz 旧名

    doc = Document(path, "pdf", "pymupdf", _ver(pymupdf))
    with pymupdf.open(path) as pdf:
        pages = min(pdf.page_count, MAX_PDF_PAGES)
        parts = []
        for i in range(pages):
            parts.append(pdf[i].get_text())
        doc.text = "\n".join(parts).strip()
        doc.meta = {"pages": pdf.page_count, "producer": pdf.metadata.get("producer", "")}
        doc.title = pdf.metadata.get("title", "")
    if not doc.text.strip():
        _ocr_fallback(doc)
    return doc


def _ocr_fallback(doc: Document) -> None:
    """无文本层扫描件 → WinRT OCR 收编（09-17 落地，原「待 OCR 阶段」兑现）。

    OCR 出字：text 回填 + parser_ver 加 +winrt-ocr1 后缀（版本参与
    缓存失效）。引擎不在位/零字（纯图印章页）：维持原失败口径。
    """
    from paistation.sense.localfiles.ocr import OCR_VER, OcrUnavailable, ocr_pdf

    try:
        text = ocr_pdf(doc.path)
    except OcrUnavailable as exc:
        raise ExtractionError(f"PDF 无文本层（OCR 不可用: {exc})") from exc
    if not text.strip():
        raise ExtractionError("PDF 无文本层且 OCR 零字（纯图/印章页）")
    doc.text = text
    doc.parser_ver = f"{doc.parser_ver}+{OCR_VER}"
    doc.meta["ocr"] = OCR_VER


def extract_docx(path: str) -> Document:
    import docx

    d = docx.Document(path)
    doc = Document(path, "word", "python-docx", _ver(docx))
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for table in d.tables:  # 表格逐行拍平（html 化的 tab 分隔）
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append("\t".join(cells))
    core = d.core_properties
    doc.title = core.title or os.path.basename(path)
    doc.meta = {"author": core.author or "", "modified": str(core.modified or "")}
    doc.text = "\n".join(parts)
    return doc


def extract_xlsx(path: str) -> Document:
    import openpyxl

    doc = Document(path, "excel", "openpyxl", _ver(openpyxl))
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    lines, cells = [], 0
    try:
        for ws in wb.worksheets:
            lines.append(f"## Sheet: {ws.title}")
            for row in ws.iter_rows(values_only=True):
                vals = ["" if v is None else str(v) for v in row]
                if any(v.strip() for v in vals):
                    lines.append("\t".join(vals))
                    cells += len(vals)
                    if cells >= MAX_EXCEL_CELLS:
                        lines.append("…（截断：超出单元格上限）")
                        break
            if cells >= MAX_EXCEL_CELLS:
                break
    finally:
        wb.close()
    doc.text = "\n".join(lines)
    doc.title = os.path.basename(path)
    return doc


def extract_pptx(path: str) -> Document:
    from pptx import Presentation

    doc = Document(path, "powerpoint", "python-pptx", _ver(
        __import__("pptx")))
    prs = Presentation(path)
    parts = []
    for i, slide in enumerate(prs.slides, start=1):
        parts.append(f"## Slide {i}")
        for shape in slide.shapes:
            if shape.has_text_frame:
                txt = shape.text_frame.text.strip()
                if txt:
                    parts.append(txt)
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    parts.append("\t".join(c.text.strip() for c in row.cells))
    doc.text = "\n".join(parts)
    doc.title = os.path.basename(path)
    return doc


def extract_msg(path: str) -> Document:
    import extract_msg

    msg = extract_msg.openMsg(path)
    try:
        doc = Document(path, "email_msg", "extract_msg", _ver(extract_msg))
        body = msg.body or msg.htmlBody or b""
        if isinstance(body, bytes):
            body = body.decode("utf-8", errors="replace")
        doc.text = (body or "").strip()
        doc.title = msg.subject or ""
        doc.meta = {
            "from": str(msg.sender or ""),
            "to": str(msg.to or ""),
            "cc": str(msg.cc or ""),
            "date": str(msg.date or ""),
        }
        doc.attachments = [a.longFilename or a.shortFilename or ""
                           for a in (msg.attachments or [])]
    finally:
        msg.close()
    return doc


def extract_eml(path: str) -> Document:
    with open(path, "rb") as fh:
        # policy=default：RFC2047 编码头（=?utf-8?b?...?=）自动解码
        mail = email.message_from_binary_file(fh, policy=email.policy.default)
    doc = Document(path, "email_eml", "stdlib-email", sys.version.split()[0])
    doc.title = str(mail.get("Subject", ""))
    doc.meta = {"from": str(mail.get("From", "")),
                "to": str(mail.get("To", "")),
                "date": str(mail.get("Date", ""))}
    bodies, atts = [], []
    for part in mail.walk():
        if part.get_content_maintype() == "multipart":
            continue
        fname = part.get_filename()
        if fname:
            atts.append(fname)
            continue
        if part.get_content_maintype() == "text":
            payload = part.get_payload(decode=True) or b""
            charset = part.get_content_charset() or "utf-8"
            bodies.append(payload.decode(charset, errors="replace"))
    doc.text = "\n".join(bodies).strip()
    doc.attachments = atts
    return doc


def extract_plaintext(path: str) -> Document:
    doc = Document(path, "text", "plaintext", "1")
    with open(path, "rb") as fh:
        raw = fh.read(MAX_TEXT_BYTES)
    doc.text = raw.decode("utf-8", errors="replace").strip("﻿\r\n ")
    doc.title = os.path.basename(path)
    return doc


def metadata_only(path: str, kind: str = "unknown") -> Document:
    """零解析路由：只登记元数据（图片/音视频/压缩包/二进制/红线外置）。"""
    st = os.stat(path)
    doc = Document(path, kind, "metadata-only", "1")
    doc.title = os.path.basename(path)
    doc.meta = {"size": st.st_size, "mtime": st.st_mtime}
    return doc


def extract_image_ocr(path: str) -> Document:
    """图片 OCR 路由：聊天截图/票据/现场照片内文字。

    零字（纯风景/人像）→ metadata 语义登记（ok 不 failed，指纹仍
    记 image-ocr——cache_hit 会命中，重扫不再重复 OCR）；引擎不在
    位 → 同样优雅降级 metadata（无语言包机器不炸批）。
    """
    from paistation.sense.localfiles.ocr import OCR_VER, OcrUnavailable, ocr_image

    doc = Document(path, "image", "image-ocr", f"1+{OCR_VER}")
    try:
        text = ocr_image(path)
    except OcrUnavailable:
        md = metadata_only(path, "image")
        md.parser_id, md.parser_ver = "image-ocr", f"1+{OCR_VER}"
        return md
    if text.strip():
        doc.text = text
    doc.title = os.path.basename(path)
    return doc


PARSERS = {
    "pymupdf": extract_pdf,
    "python-docx": extract_docx,
    "openpyxl": extract_xlsx,
    "python-pptx": extract_pptx,
    "extract_msg": extract_msg,
    "stdlib-email": extract_eml,
    "plaintext": extract_plaintext,
    "image-ocr": extract_image_ocr,
    "metadata-only": lambda p: metadata_only(p),
}


def extract(path: str, kind: str, parser_id: str) -> Document:
    """路由分发 + 统一异常包装。"""
    parser = PARSERS.get(parser_id)
    if parser is None:
        return metadata_only(path, kind)
    try:
        doc = parser(path)
        doc.kind = kind if kind != "unknown" else doc.kind
        return doc
    except ExtractionError:
        raise
    except Exception as exc:  # 毒文件常态化：单文件失败绝不崩批
        raise ExtractionError(f"{parser_id} 解析失败 {path}: {exc}") from exc
