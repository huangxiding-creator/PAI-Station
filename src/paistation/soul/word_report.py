"""三模型画像 Word 报告（提案 1.6 支柱三/后续版本项）。

三模型视角：事实（一级直写）/ 行为（感知推断）/ 成长（指数轨迹）。
python-docx 输出，中文字体微软雅黑（复用 hundun_to_docx 实战配方）。
"""
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

_FONT = "微软雅黑"


def _set_font(run) -> None:
    run.font.name = _FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), _FONT)


def _head(doc, text: str, level: int = 1) -> None:
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        _set_font(run)


def _para(doc, text: str, bold: bool = False, center: bool = False) -> None:
    p = doc.add_paragraph()
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.bold = bold
    _set_font(run)


def profile_word_report(facts: list, title: str = "三模型画像报告") -> Document:
    """facts [{key, value, source}] → 三视角 Word 文档。"""
    doc = Document()
    _para(doc, title, bold=True, center=True)
    _head(doc, "一、事实模型（一级直写）", 1)
    hard = [f for f in facts if f.get("source") in ("名片", "档案", "登记")]
    for f in hard or facts[:1]:
        _para(doc, f"{f['key']}：{f['value']}")
    if not facts:
        _para(doc, "（暂无画像数据——诚实优先，不编造）")

    _head(doc, "二、行为模型（感知推断）", 1)
    behav = [f for f in facts if f.get("source") in ("感知", "行为")]
    for f in behav or []:
        _para(doc, f"{f['key']}：{f['value']}")
    if not behav and facts:
        _para(doc, "（行为样本积累中）")

    _head(doc, "三、成长模型（指数轨迹）", 1)
    growth = [f for f in facts if f.get("source") in ("growth", "成长")]
    for f in growth or []:
        _para(doc, f"{f['key']}：{f['value']}")
    if not growth and facts:
        _para(doc, "（成长指数首月后生成）")
    return doc


def save_docx(doc: Document, path: str) -> None:
    doc.save(path)
