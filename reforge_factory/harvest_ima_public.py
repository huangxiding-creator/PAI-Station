# -*- coding: utf-8 -*-
"""harvest_ima_public — 腾讯 ima 公开知识库收割腿 (EPC-BID26 · ima_public 渠道).

链路 (驱动 ResearchFactory-Eng/EPC100/collectors/ima_public, 不重复造轮):
  ImaOpenApi.search_knowledge_base (面词扇出, 单核心词>组合词 0926 实测)
  → 库名域词门 (招标|投标|评标|工程|EPC|采购|承包|造价|FIDIC|BIM|供应链)
  → search_knowledge 库内搜 (面词, 标题再过域词门)
  → get_media_info 签名直链 (res-skb + X-IMA-* 头, 直连关代理)
  → 下载 → 文本抽取 (pdf=pymupdf / docx=python-docx / html=剥标签 / txt=原样)
  → md 证据件 (>300字门, 非 SERP 形) 落 data/harvest/ima_public/<cid>/

纪律: 不 import ammo_pool 不 ingest (本腿只产文件); 账号安全四件套走
collector.Throttle (3s/req + 日限额300 + 冷却300s + 熔断5连败); 断点
state 存本腿 out 目录 (不碰 collector 既有断点/目录卡); .doc/.xlsx 等无
抽取器件如实跳过不凑数; 限流/熔断/零结果如实报。

用法: py -3.11 reforge_factory/harvest_ima_public.py --campaign EPC-BID26 [--dry]
"""
from __future__ import annotations

import argparse
import hashlib
import io
import json
import re
import sys
import time
import urllib.request
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

REPO = Path(__file__).resolve().parents[1]                  # E:\AI-Station
IMA_PKG = REPO / "ResearchFactory-Eng" / "EPC100" / "collectors"
sys.path.insert(0, str(IMA_PKG))
from ima_public.client import ImaApiError, ImaOpenApi        # noqa: E402
from ima_public.discover import snapshot_catalog             # noqa: E402
from ima_public.harvest import safe_filename                 # noqa: E402
from ima_public.throttle import CircuitOpen, DailyLimitExceeded, Throttle  # noqa: E402

ENGINE = "ima_public:kb"          # 引擎标签 (含渠道 id 子串)
# (面词查询, 目标 EEI) — 四面优先: 国际对标Q23/Q24 → 技术Q18/Q19/Q20 →
# 组织人事Q21/Q22 → 渠道生态Q25/Q26; 域词兜底 Q1-E1
PLAN: list[tuple[str, str]] = [
    ("FIDIC", "Q23-E1"),
    ("国际工程", "Q24-E1"),
    ("对外承包", "Q24-E1"),
    ("电子招投标", "Q19-E1"),
    ("BIM", "Q19-E1"),
    ("评标", "Q21-E1"),
    ("造价", "Q21-E1"),
    ("招标代理", "Q25-E1"),
    ("公共资源交易", "Q26-E1"),
    ("供应链", "Q25-E1"),
    ("采购", "Q25-E1"),
    ("总承包", "Q1-E1"),
    ("招标投标", "Q1-E1"),
]
DOMAIN_RE = re.compile(
    r"招标|投标|评标|工程|EPC|采购|承包|造价|FIDIC|BIM|供应链")
MAX_KBS = 7                 # 收割库数上限
QUERIES_PER_KB = 2          # 每库库内搜查询数上限
KEEP_PER_KB = 3             # 每库存件上限 (质量优先)
TARGET_KEEP = 12            # 总存件目标 (契约 5-15)
MIN_CHARS = 300             # 证据件字数门 (>300字)
FETCH_TEXTY = (".pdf", ".docx", ".txt", ".md", ".html")   # 有抽取器的件


# ---------------------------------------------------------------- 抽取器
def _pdf_text(body: bytes) -> str:
    import fitz                                              # pymupdf
    with fitz.open(stream=body, filetype="pdf") as doc:
        return "\n".join(p.get_text() for p in doc)


def _docx_text(body: bytes) -> str:
    import docx
    d = docx.Document(io.BytesIO(body))
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:                                       # 表格文字也收
        for row in t.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    return "\n".join(parts)


_TAG_RE = re.compile(r"<(script|style)[^>]*>.*?</\1>|<[^>]+>",
                     re.S | re.I)


def _html_text(body: bytes) -> str:
    txt = _TAG_RE.sub(" ",
                      body.decode("utf-8", errors="replace"))
    return re.sub(r"\n{3,}| {2,}", "\n", txt).strip()


def extract_text(body: bytes, ext: str) -> str:
    """按扩展名抽正文; 无抽取器/抽取失败返回空 (如实跳)."""
    try:
        if ext == ".pdf":
            return _pdf_text(body)
        if ext == ".docx":
            return _docx_text(body)
        if ext == ".html":
            return _html_text(body)
        if ext in (".txt", ".md"):
            return body.decode("utf-8", errors="replace")
    except Exception:                                        # noqa: BLE001
        return ""
    return ""


def _download(opener, url: str, headers: dict) -> bytes:
    req = urllib.request.Request(url, headers={
        "User-Agent": "Mozilla/5.0", **(headers or {})})
    with opener.open(req, timeout=120) as resp:
        return resp.read()


# ---------------------------------------------------------------- 主腿
def run(cid: str, dry: bool = False) -> int:
    out_dir = REPO / "data" / "harvest" / "ima_public" / cid
    out_dir.mkdir(parents=True, exist_ok=True)
    state_path = out_dir / "_state_ima_public.json"
    state: dict = {}
    if state_path.is_file():                                 # 断点续跑
        try:
            state = json.loads(state_path.read_text(encoding="utf-8"))
        except ValueError:
            state = {}
    records: list[dict] = list(state.get("records", []))
    done_ids: set[str] = set(state.get("done_ids", []))
    text_hashes: set[str] = set(state.get("text_hashes", []))

    try:
        api = ImaOpenApi()
    except ImaApiError as exc:
        print(f"[ima_public] 凭证缺位, 收兵: {exc}")
        return 2
    throttle = Throttle()
    opener = urllib.request.build_opener(
        urllib.request.ProxyHandler({}))                     # 直连 (y869)

    # ---- 腿1: 面词扇出搜库 → 库名域词门
    kb_queries: dict[str, list[tuple[str, str]]] = {}        # kb_id → [(q,eei)]
    kb_names: dict[str, str] = {}
    kb_counts: dict[str, int] = {}
    for q, eei in PLAN:
        throttle.pace()
        try:
            d = api.search_knowledge_base(query=q, limit=8)
        except ImaApiError as exc:
            throttle.record_failure()
            print(f"[ima_public] search '{q}' 失败: {str(exc)[:90]}")
            continue
        throttle.record_ok()
        for it in (d.get("info_list") or []):
            kid, name = it.get("kb_id", ""), it.get("kb_name", "") or ""
            if not kid:
                continue
            if not DOMAIN_RE.search(name):                   # 库名域词门
                continue
            kb_queries.setdefault(kid, []).append((q, eei))
            kb_names[kid] = name
            kb_counts[kid] = max(kb_counts.get(kid, 0),
                                 int(it.get("content_count", 0) or 0))
        time.sleep(0.2)
    print(f"[ima_public] discover 命中库 {len(kb_queries)} "
          f"(queries={len(PLAN)})")
    # 目录卡快照 (本腿私有, 不碰 EPC100 共账)
    from ima_public.discover import KBCard
    snapshot_catalog([KBCard(kb_id=k, name=kb_names[k],
                             content_count=0, member_count=0)
                      for k in kb_queries], out_dir / "_catalog.json")

    # 排序: 首命中查询的 PLAN 位次 (四面优先序) → 命中面词数 → 内容量;
    # 每查询至多 2 库 (防单面词库海淹没四面多样性)
    q_rank = {q: i for i, (q, _) in enumerate(PLAN)}
    ranked_all = sorted(
        kb_queries,
        key=lambda k: (min(q_rank[q] for q, _ in kb_queries[k]),
                       -len(kb_queries[k]), -kb_counts.get(k, 0), k))
    ranked, per_q = [], {}
    for k in ranked_all:
        key = min((q_rank[q], q) for q, _ in kb_queries[k])[1]
        if per_q.get(key, 0) >= 2:
            continue
        per_q[key] = per_q.get(key, 0) + 1
        ranked.append(k)
        if len(ranked) >= MAX_KBS:
            break
    if dry:
        for k in ranked:
            print(f"  [dry] {kb_names[k][:36]} ({kb_counts.get(k, 0)}条) ← "
                  f"{[q for q, _ in kb_queries[k]]}")
        print(f"[ima_public] DRY: 拟收 {len(ranked)} 库")
        return 0

    n_keep = n_skip_gate = n_skip_ext = n_skip_thin = n_fail = 0
    blockers: list[str] = []
    try:
        for kid in ranked[:MAX_KBS]:
            if n_keep >= TARGET_KEEP:
                break
            kname = kb_names[kid]
            kept_kb = 0
            # ---- 腿2: 库内搜 (面词), 标题域词门
            cand: dict[str, tuple[str, str, str]] = {}       # mid → (q,eei,title)
            for q, eei in kb_queries[kid][:QUERIES_PER_KB]:
                throttle.pace()
                try:
                    d = api.search_knowledge(query=q, kb_id=kid, limit=8)
                except ImaApiError as exc:
                    throttle.record_failure()
                    print(f"[ima_public] in-KB '{q}' @{kname[:20]} 失败: "
                          f"{str(exc)[:80]}")
                    continue
                throttle.record_ok()
                for it in (d.get("info_list") or []):
                    mid, title = it.get("media_id", ""), \
                        (it.get("title") or "").strip()
                    if not mid or mid in done_ids:
                        continue
                    if mid.startswith("folder_"):
                        continue              # 目录条目: get_media_info 必败
                    if not DOMAIN_RE.search(title):          # 标题域词门
                        n_skip_gate += 1
                        continue
                    cand.setdefault(mid, (q, eei, title))
            print(f"[ima_public] 库 '{kname[:32]}' 候选 {len(cand)} 条")
            # ---- 腿3: 逐件签名直链下载 → 抽取 → md 存件
            for mid, (q, eei, title) in cand.items():
                if kept_kb >= KEEP_PER_KB or n_keep >= TARGET_KEEP:
                    break
                throttle.pace()
                try:
                    info = api.get_media_info(mid)
                except ImaApiError:
                    throttle.record_failure()
                    n_fail += 1
                    continue
                throttle.record_ok()
                ui = info.get("url_info") or {}
                url, heads = ui.get("url", ""), ui.get("headers") or {}
                if not url.startswith("http"):
                    # ima 笔记/网链类无文件直链: 留痕跳过, 不计失败不耗熔断
                    n_skip_ext += 1
                    done_ids.add(mid)
                    continue
                ext = Path(url.split("?")[0]).suffix.lower() \
                    or Path(title).suffix.lower()
                if ext == ".doc":                             # 老格式无抽取器
                    n_skip_ext += 1
                    done_ids.add(mid)
                    continue
                if ext not in FETCH_TEXTY:
                    n_skip_ext += 1
                    done_ids.add(mid)
                    continue
                try:
                    body = _download(opener, url, heads)
                except Exception:                             # noqa: BLE001
                    n_fail += 1
                    time.sleep(1.0)
                    continue
                time.sleep(1.0)                               # 下载腿礼貌间隔
                text = extract_text(body, ext).strip()
                if len(text) < MIN_CHARS:                     # 薄件/扫描件如实跳
                    n_skip_thin += 1
                    done_ids.add(mid)
                    continue
                th = hashlib.sha1(text.encode("utf-8")).hexdigest()
                if th in text_hashes:                         # 同文异名去重
                    n_skip_thin += 1
                    done_ids.add(mid)
                    continue
                text_hashes.add(th)
                fp = out_dir / f"{safe_filename(title)[:64]}.md"
                head = (f"# {title}\n\n> 来源: 腾讯ima公开知识库 "
                        f"「{kname}」(库内搜 '{q}')\n> 渠道: ima_public "
                        f"(ima.qq.com OpenAPI)\n\n")
                fp.write_text(head + text, encoding="utf-8")
                records.append({"file": str(fp), "url": "",
                                "engine": ENGINE, "tree": eei,
                                "cred": "research"})
                done_ids.add(mid)
                kept_kb += 1
                n_keep += 1
                print(f"  [keep] {eei} {title[:44]} "
                      f"({len(text)}字)")
                state = {"records": records, "text_hashes": sorted(text_hashes),
                         "done_ids": sorted(done_ids)}
                state_path.write_text(
                    json.dumps(state, ensure_ascii=False, indent=1),
                    encoding="utf-8")                         # 每件落账断点
    except DailyLimitExceeded as exc:
        blockers.append(f"daily_limit: {exc}")
    except CircuitOpen as exc:
        blockers.append(f"circuit: {exc}")

    # ---- 收尾: 清单 + 汇总
    man = out_dir / "_records.jsonl"
    man.write_text("\n".join(json.dumps(r, ensure_ascii=False)
                             for r in records) + "\n", encoding="utf-8")
    state = {"records": records, "done_ids": sorted(done_ids),
             "text_hashes": sorted(text_hashes), "blockers": blockers}
    state_path.write_text(json.dumps(state, ensure_ascii=False, indent=1),
                          encoding="utf-8")
    print(f"[ima_public] APPLY: 存件 {n_keep} | 域门跳 {n_skip_gate} | "
          f"无抽取器跳 {n_skip_ext} | 薄件跳 {n_skip_thin} | 失败 {n_fail}"
          + (f" | blockers: {blockers}" if blockers else ""))
    print(f"[ima_public] 清单 → {man}")
    return 0


def main() -> int:
    ap = argparse.ArgumentParser(
        description="ima 公开知识库收割腿 (ima_public 渠道, 只产文件不入池)")
    ap.add_argument("--campaign", default="EPC-BID26")
    ap.add_argument("--dry", action="store_true")
    ns = ap.parse_args()
    return run(ns.campaign, ns.dry)


if __name__ == "__main__":
    sys.exit(main())
